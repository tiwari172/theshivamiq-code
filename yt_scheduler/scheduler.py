"""theshivamiq - YouTube Shorts scheduler.

Excel ki "Post Schedule" sheet padh kar videos ko YouTube par private + scheduled upload karta hai.
Sirf tab chalta hai jab aap khud chalao (koi background loop nahi). Commands ke liye SETUP.md dekho.
"""

from __future__ import annotations

import argparse
import http.client
import json
import logging
import os
import random
import shutil
import sys
import time
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import httplib2
from dotenv import load_dotenv
from google.auth.exceptions import RefreshError
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from googleapiclient.http import MediaFileUpload
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

SCRIPT_DIR = Path(__file__).resolve().parent
SCOPES = ["https://www.googleapis.com/auth/youtube.upload"]
HARD_DAILY_CAP = 6                      # 10,000 units / 1,600 per upload = 6
CHUNK_SIZE = 4 * 1024 * 1024            # 4 MB
MIN_LEAD_TIME = timedelta(minutes=30)   # publishAt kam se kam itna aage ho
RETRYABLE_STATUS = {408, 429, 500, 502, 503, 504}
RETRYABLE_EXCEPTIONS = (httplib2.HttpLib2Error, http.client.HTTPException, OSError)
QUOTA_REASONS = {"quotaExceeded", "dailyLimitExceeded", "uploadLimitExceeded"}
DATE_FORMATS = ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M")
NEEDED_COLUMNS = (
    "video_filename", "yt_title", "yt_description", "yt_tags", "schedule_datetime",
    "target_platform", "status", "log_message", "youtube_video_id", "processed_at",
)
MAX_BACKUPS = 20
QUOTA_TZ = ZoneInfo("America/Los_Angeles")   # YouTube ka daily quota Pacific Time ki aadhi raat par reset hota hai
LOCK_STALE = timedelta(hours=6)              # isse purani lock file = pichhla run crash hua tha

log = logging.getLogger("yt_scheduler")


class UploadError(Exception):
    """Upload fail hua; message Excel ke log_message mein jaata hai."""


class QuotaExceeded(Exception):
    """YouTube ne bola aaj ka quota khatam."""


@dataclass
class Config:
    """.env se padhi gayi saari settings."""

    video_dir: Path
    excel_path: Path
    sheet_name: str
    tz: ZoneInfo
    client_secrets: Path
    token_file: Path
    category_id: str
    daily_limit: int
    min_sleep: float
    max_sleep: float
    max_retries: int
    quota_file: Path
    pending_file: Path
    log_file: Path
    backup_dir: Path
    synthetic_media: bool | None = None
    lock_file: Path | None = None


@dataclass
class Row:
    """Excel ki ek data row (sirf jo kaam ki cheezein hain)."""

    excel_row: int
    case_no: str
    part: str
    filename: str
    title: str
    description: str
    tags: str
    when_raw: Any
    platform: str
    status: str

    @property
    def label(self) -> str:
        """Log ke liye chhota naam: row number + file ka naam."""
        return f"Row {self.excel_row} ({self.filename or 'no file'})"


# ---------------------------------------------------------------- config / logging

def zone(name: str) -> ZoneInfo:
    """TIMEZONE ka naam -> ZoneInfo; galat naam ho to saaf message."""
    try:
        return ZoneInfo(name)
    except Exception:
        raise SystemExit(f".env mein TIMEZONE={name} samajh nahi aaya (Asia/Kolkata likho). "
                         "Ya '0 INSTALL.bat' dobara chalao (tzdata).") from None


def resolve_path(value: str) -> Path:
    """Relative path ko script folder ke hisaab se poora path banata hai."""
    p = Path(value.strip().strip('"')).expanduser()
    return p if p.is_absolute() else SCRIPT_DIR / p


def env_num(name: str, default: str, kind: type = int) -> Any:
    """.env ka number; khaali ya galat ho to saaf message (traceback nahi)."""
    raw = (os.getenv(name) or "").strip() or default
    try:
        return kind(raw)
    except ValueError:
        raise SystemExit(f".env mein {name}={raw} number nahi hai. Sirf number likho (jaise {default}).") from None


def env_yes_no(name: str) -> bool | None:
    """Haan/Yes -> True, Nahi/No -> False, khaali -> None (YouTube ko kuch nahi bhejna)."""
    raw = (os.getenv(name) or "").strip().lower()
    if raw in ("haan", "han", "yes", "y", "true", "1"):
        return True
    if raw in ("nahi", "nahin", "no", "n", "false", "0"):
        return False
    return None


def load_config() -> Config:
    """.env file (script ke bagal wali) se settings padhta hai."""
    load_dotenv(SCRIPT_DIR / ".env")
    env = lambda name, default: (os.getenv(name) or "").strip() or default  # noqa: E731  (khaali = default)
    limit = env_num("DAILY_UPLOAD_LIMIT", "6")
    return Config(
        video_dir=resolve_path(env("VIDEO_DIR", "output")),
        excel_path=resolve_path(env("EXCEL_PATH", "schedule.xlsx")),
        sheet_name=env("SHEET_NAME", "Post Schedule"),
        tz=zone(env("TIMEZONE", "Asia/Kolkata")),
        client_secrets=resolve_path(env("YT_CLIENT_SECRETS_FILE", "client_secret.json")),
        token_file=resolve_path(env("YT_TOKEN_FILE", "token.json")),
        category_id=env("YT_CATEGORY_ID", "24"),
        daily_limit=max(0, min(limit, HARD_DAILY_CAP)),
        min_sleep=env_num("MIN_SLEEP_SECONDS", "5", float),
        max_sleep=env_num("MAX_SLEEP_SECONDS", "15", float),
        max_retries=env_num("MAX_RETRIES", "5"),
        quota_file=SCRIPT_DIR / "quota_state.json",
        pending_file=SCRIPT_DIR / "pending_updates.json",
        log_file=SCRIPT_DIR / "scheduler.log",
        backup_dir=SCRIPT_DIR / "backups",
        synthetic_media=env_yes_no("YT_CONTAINS_SYNTHETIC_MEDIA"),
        lock_file=SCRIPT_DIR / "scheduler.lock",
    )


def setup_logging(log_file: Path) -> None:
    """Console aur scheduler.log dono mein log likhne ka intezaam karta hai."""
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    log.setLevel(logging.INFO)
    file_handler = logging.FileHandler(log_file, encoding="utf-8")
    file_handler.setFormatter(logging.Formatter("%(asctime)s | %(levelname)-7s | %(message)s", "%Y-%m-%d %H:%M:%S"))
    console = logging.StreamHandler(sys.stdout)
    console.setFormatter(logging.Formatter("%(message)s"))
    log.addHandler(file_handler)
    log.addHandler(console)


# ---------------------------------------------------------------- quota

def utc_today() -> str:
    """Quota ka "aaj" - Pacific Time ki date (YouTube quota PT midnight par reset hota hai,
    yaani India mein dopahar ~12:30 / 1:30). Naam purana hi rakha hai."""
    return datetime.now(QUOTA_TZ).strftime("%Y-%m-%d")


def load_quota(cfg: Config) -> dict[str, Any]:
    """quota_state.json padhta hai; date badal gayi ho to counter 0 kar deta hai."""
    state: dict[str, Any] = {}
    if cfg.quota_file.exists():
        try:
            state = json.loads(cfg.quota_file.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            log.warning("quota_state.json padh nahi paaya, naya bana raha hun.")
    if state.get("date") != utc_today():
        state = {"date": utc_today(), "uploads": 0}
    return state


def save_quota(cfg: Config, state: dict[str, Any]) -> None:
    """quota_state.json mein aaj ka counter likhta hai."""
    cfg.quota_file.write_text(json.dumps(state, indent=2), encoding="utf-8")


def quota_left(cfg: Config, state: dict[str, Any]) -> int:
    """Aaj kitne upload aur ho sakte hain."""
    return max(0, cfg.daily_limit - int(state.get("uploads", 0)))


# ---------------------------------------------------------------- ek saath do run nahi

def take_lock(cfg: Config) -> None:
    """scheduler.lock banata hai; doosra run chal raha ho to ruk jaata hai (warna same video 2 baar upload)."""
    lock = cfg.lock_file
    if lock is None:
        return
    for _ in range(2):
        try:
            fd = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
        except FileExistsError:
            try:
                age = datetime.now() - datetime.fromtimestamp(lock.stat().st_mtime)
            except OSError:
                continue
            if age > LOCK_STALE:
                log.warning("Purani scheduler.lock mili (pichhla run beech mein band hua tha) - hata raha hun.")
                lock.unlink(missing_ok=True)
                continue
            raise SystemExit("Scheduler pehle se chal raha hai (doosri black window dekho). Wo khatam hone do. "
                             "Agar koi window nahi chal rahi to yt_scheduler\\scheduler.lock delete karo.")
        with os.fdopen(fd, "w") as f:
            f.write(f"{os.getpid()} {datetime.now():%Y-%m-%d %H:%M:%S}\n")
        return


def release_lock(cfg: Config) -> None:
    """scheduler.lock hata deta hai."""
    if cfg.lock_file is not None:
        cfg.lock_file.unlink(missing_ok=True)


# ---------------------------------------------------------------- excel helpers

def excel_is_locked(path: Path) -> bool:
    """True agar file MS Excel mein khuli hai (Windows write-lock laga deta hai)."""
    try:
        with open(path, "r+b"):
            return False
    except PermissionError:
        return True


def open_workbook(cfg: Config, read_only: bool) -> tuple[Workbook, Worksheet]:
    """Workbook kholta hai aur "Post Schedule" sheet deta hai."""
    if not cfg.excel_path.exists():
        raise SystemExit(f"Excel file nahi mili: {cfg.excel_path}  (.env mein EXCEL_PATH check karo)")
    keep_vba = cfg.excel_path.suffix.lower() == ".xlsm"
    wb = load_workbook(cfg.excel_path, read_only=read_only, keep_vba=keep_vba)
    if cfg.sheet_name not in wb.sheetnames:
        raise SystemExit(f'Excel mein "{cfg.sheet_name}" sheet nahi mili. Sheets: {", ".join(wb.sheetnames)}')
    return wb, wb[cfg.sheet_name]


def header_map(ws: Worksheet) -> dict[str, int]:
    """Row 1 ke header -> column number (1 se shuru) ka map banata hai."""
    first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    cols = {str(v).strip().lower(): i + 1 for i, v in enumerate(first) if v is not None and str(v).strip()}
    missing = [c for c in NEEDED_COLUMNS if c not in cols]
    if missing:
        raise SystemExit(f"Post Schedule sheet mein ye columns nahi mile: {', '.join(missing)}")
    return cols


def text(value: Any) -> str:
    """Cell ki value ko saaf text banata hai (None -> "")."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_rows(ws: Worksheet, cols: dict[str, int]) -> list[Row]:
    """Sheet ki saari bhari hui rows padhta hai."""
    rows: list[Row] = []
    for n, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not values or all(v is None or text(v) == "" for v in values):
            continue
        get = lambda name: values[cols[name] - 1] if name in cols and cols[name] - 1 < len(values) else None  # noqa: E731
        rows.append(Row(
            excel_row=n,
            case_no=text(get("case_no")),
            part=text(get("part")),
            filename=text(get("video_filename")),
            title=text(get("yt_title")),
            description=text(get("yt_description")),
            tags=text(get("yt_tags")),
            when_raw=get("schedule_datetime"),
            platform=text(get("target_platform")).upper(),
            status=text(get("status")).upper(),
        ))
    return rows


def write_cells(ws: Worksheet, cols: dict[str, int], excel_row: int, updates: dict[str, str]) -> None:
    """Ek row ke diye gaye columns mein nayi value likhta hai (baaki kuch nahi chhuta)."""
    for name, value in updates.items():
        ws.cell(row=excel_row, column=cols[name]).value = value


def backup_excel(cfg: Config) -> None:
    """Excel mein likhne se pehle backups\\ folder mein ek copy rakhta hai (purani 20 hi rakhta hai)."""
    cfg.backup_dir.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    dest = cfg.backup_dir / f"{cfg.excel_path.stem}_{stamp}{cfg.excel_path.suffix}"
    shutil.copy2(cfg.excel_path, dest)
    old = sorted(cfg.backup_dir.glob(f"{cfg.excel_path.stem}_*{cfg.excel_path.suffix}"))
    for f in old[:-MAX_BACKUPS]:
        f.unlink(missing_ok=True)
    log.info(f"Excel ka backup: {dest.name}")


def load_pending(cfg: Config) -> list[dict[str, Any]]:
    """pending_updates.json (jo pichhli baar save nahi ho paaye) padhta hai."""
    if not cfg.pending_file.exists():
        return []
    try:
        return json.loads(cfg.pending_file.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        log.error("pending_updates.json kharab hai - use haath se dekho, delete mat karo.")
        raise SystemExit(1)


def add_pending(cfg: Config, excel_row: int, filename: str, updates: dict[str, str]) -> None:
    """Save fail hua to update ko pending_updates.json mein jodta hai taaki kuch khoye nahi."""
    pending = load_pending(cfg)
    pending.append({"row": excel_row, "video_filename": filename, "updates": updates})
    cfg.pending_file.write_text(json.dumps(pending, indent=2, ensure_ascii=False), encoding="utf-8")


def find_pending_row(ws: Worksheet, cols: dict[str, int], item: dict[str, Any]) -> int | None:
    """Pending update kis row ka hai: pehle row number, file naam na mile to naam se dhoondta hai."""
    col = cols["video_filename"]
    if text(ws.cell(row=item["row"], column=col).value) == item["video_filename"]:
        return int(item["row"])
    for n in range(2, ws.max_row + 1):
        if text(ws.cell(row=n, column=col).value) == item["video_filename"]:
            return n
    return None


def apply_pending(wb: Workbook, ws: Worksheet, cols: dict[str, int], cfg: Config) -> None:
    """Pichhli baar ke bache updates Excel mein daal kar save karta hai."""
    pending = load_pending(cfg)
    if not pending:
        return
    log.info(f"pending_updates.json se {len(pending)} purane update Excel mein daal raha hun...")
    for item in pending:
        n = find_pending_row(ws, cols, item)
        if n is None:
            log.error(f"Pending update ki row nahi mili ({item['video_filename']}) - pending_updates.json haath se dekho.")
            raise SystemExit(1)
        write_cells(ws, cols, n, item["updates"])
    wb.save(cfg.excel_path)
    cfg.pending_file.rename(cfg.pending_file.with_name(f"pending_applied_{datetime.now():%Y%m%d_%H%M%S}.json"))
    log.info("Purane updates Excel mein save ho gaye.")


def overlay_pending(rows: list[Row], cfg: Config) -> None:
    """Dry-run/status ke liye: pending updates ka status rows par (sirf memory mein) laga deta hai."""
    for item in load_pending(cfg):
        for r in rows:
            if r.filename == item["video_filename"] and "status" in item["updates"]:
                r.status = item["updates"]["status"].upper()


def save_row(wb: Workbook, ws: Worksheet, cols: dict[str, int], cfg: Config, row: Row, updates: dict[str, str]) -> None:
    """Ek row update karke turant Excel save karta hai; save fail ho to pending file mein rakhta hai."""
    write_cells(ws, cols, row.excel_row, updates)
    try:
        wb.save(cfg.excel_path)
    except OSError as e:
        add_pending(cfg, row.excel_row, row.filename, updates)
        log.error(f"Excel save nahi hui ({e.__class__.__name__}). Update pending_updates.json mein rakh diya - "
                  "agli baar script chalane par apne aap Excel mein chala jayega.")


# ---------------------------------------------------------------- validation

def now_ist(cfg: Config) -> datetime:
    """Abhi ka time, IST mein."""
    return datetime.now(cfg.tz)


def parse_schedule(value: Any, cfg: Config) -> datetime | None:
    """schedule_datetime (text ya Excel date) ko IST wala datetime banata hai."""
    if isinstance(value, datetime):
        return value.replace(tzinfo=cfg.tz) if value.tzinfo is None else value
    raw = text(value)
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(raw, fmt).replace(tzinfo=cfg.tz)
        except ValueError:
            continue
    return None


def is_candidate(row: Row) -> bool:
    """Row PENDING/khaali hai aur YouTube ke liye hai?"""
    return row.status in ("", "PENDING") and ("YOUTUBE" in row.platform or "BOTH" in row.platform)


def check_row(row: Row, cfg: Config) -> tuple[datetime | None, str | None]:
    """Upload se pehle saare check; (time, None) ya (None, galti ki wajah) deta hai."""
    if not row.filename:
        return None, "video_filename khaali hai"
    path = cfg.video_dir / row.filename
    if path.suffix.lower() != ".mp4":
        return None, f"file .mp4 nahi hai: {row.filename}"
    if not path.is_file():
        return None, f"video file nahi mili: {path}"
    if path.stat().st_size == 0:
        return None, "video file khaali hai (0 bytes)"
    if not row.title:
        return None, "yt_title khaali hai"
    when = parse_schedule(row.when_raw, cfg)
    if when is None:
        return None, f"schedule_datetime samajh nahi aaya: '{text(row.when_raw)}' (format: 2026-10-14 08:00:00)"
    if when < now_ist(cfg) + MIN_LEAD_TIME:
        return None, f"schedule time {when:%Y-%m-%d %H:%M} abhi se 30 min se kam aage hai / nikal gaya"
    return when, None


def split_rows(rows: list[Row], cfg: Config) -> tuple[list[tuple[Row, datetime]], list[tuple[Row, str]]]:
    """Candidate rows ko 'ready' (time ke hisaab se sorted) aur 'problem' mein baantta hai."""
    ready: list[tuple[Row, datetime]] = []
    problems: list[tuple[Row, str]] = []
    for r in rows:
        if not is_candidate(r):
            continue
        when, err = check_row(r, cfg)
        if err:
            problems.append((r, err))
        else:
            ready.append((r, when))  # type: ignore[arg-type]
    ready.sort(key=lambda x: x[1])
    return ready, problems


# ---------------------------------------------------------------- youtube

def get_youtube(cfg: Config) -> Any:
    """OAuth login (token.json reuse/refresh, warna browser) karke YouTube client deta hai."""
    creds: Credentials | None = None
    if cfg.token_file.exists():
        try:
            creds = Credentials.from_authorized_user_file(str(cfg.token_file), SCOPES)
        except (ValueError, json.JSONDecodeError):
            log.warning("token.json kharab hai, dobara login karwa raha hun.")
    if creds and creds.expired and creds.refresh_token:
        try:
            creds.refresh(Request())
        except RefreshError:
            log.warning("Token refresh nahi hua (invalid_grant). Browser mein dobara login karo.")
            creds = None
    if not creds or not creds.valid:
        if not cfg.client_secrets.exists():
            raise SystemExit(f"client_secret.json nahi mili: {cfg.client_secrets}  (SETUP.md step 3 dekho)")
        log.info("Browser khul raha hai - wahi Google account chuno jiska YouTube channel hai.")
        flow = InstalledAppFlow.from_client_secrets_file(str(cfg.client_secrets), SCOPES)
        creds = flow.run_local_server(port=0)
    cfg.token_file.write_text(creds.to_json(), encoding="utf-8")
    return build("youtube", "v3", credentials=creds, cache_discovery=False)


def clean(s: str) -> str:
    """YouTube title/description mein mana '<' '>' hata deta hai."""
    return s.replace("<", "").replace(">", "").strip()


def build_tags(raw: str) -> list[str]:
    """Comma wale tags ko list banata hai (max 15, kul 450 akshar tak)."""
    tags: list[str] = []
    total = 0
    for t in (clean(x) for x in raw.split(",")):
        if not t or len(tags) >= 15 or total + len(t) > 450:
            continue
        tags.append(t)
        total += len(t) + 1
    return tags


def to_utc_iso(when: datetime) -> str:
    """IST time ko YouTube wale UTC ISO format ("2026-10-14T02:30:00Z") mein badalta hai."""
    return when.astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def build_body(row: Row, when: datetime, cfg: Config) -> dict[str, Any]:
    """videos.insert ke liye snippet + status body banata hai."""
    description = clean(row.description)
    if "#shorts" not in description.lower():
        description = f"{description}\n\n#Shorts".strip()
    return {
        "snippet": {
            "title": clean(row.title)[:100],
            "description": description[:5000],
            "tags": build_tags(row.tags),
            "categoryId": cfg.category_id,
        },
        "status": {
            "privacyStatus": "private",
            "publishAt": to_utc_iso(when),
            "selfDeclaredMadeForKids": False,
            **({"containsSyntheticMedia": cfg.synthetic_media} if cfg.synthetic_media is not None else {}),
        },
    }


def error_reason(e: HttpError) -> tuple[str, str]:
    """HttpError se (reason, message) nikalta hai, jaise ("quotaExceeded", "...")."""
    try:
        data = json.loads(e.content.decode("utf-8"))
        err = data.get("error", {})
        reason = (err.get("errors") or [{}])[0].get("reason", "")
        return reason, err.get("message", "")
    except (ValueError, AttributeError):
        return "", str(e)


def backoff(attempt: int, max_retries: int, why: str) -> int:
    """Retry se pehle 2s, 4s, 8s... (max 60s) + jitter rukta hai; try khatam ho to error deta hai."""
    attempt += 1
    if attempt > max_retries:
        raise UploadError(f"{max_retries} baar try kiya, phir bhi fail: {why}")
    delay = min(60.0, 2.0 ** attempt) + random.uniform(0, 1)
    log.warning(f"   {why} - {delay:.0f}s baad dobara try ({attempt}/{max_retries})...")
    time.sleep(delay)
    return attempt


def upload_video(youtube: Any, path: Path, body: dict[str, Any], max_retries: int) -> str:
    """Resumable 4MB chunks mein upload karta hai, net toote to wahin se aage; video id deta hai."""
    media = MediaFileUpload(str(path), mimetype="video/mp4", chunksize=CHUNK_SIZE, resumable=True)
    request = youtube.videos().insert(part="snippet,status", body=body, media_body=media)
    response = None
    attempt = 0
    while response is None:
        try:
            status, response = request.next_chunk()
            attempt = 0
            if status:
                log.info(f"   upload {int(status.progress() * 100)}%")
        except HttpError as e:
            code = e.resp.status
            reason, message = error_reason(e)
            if reason in QUOTA_REASONS:
                raise QuotaExceeded(reason) from e
            if code in RETRYABLE_STATUS:
                attempt = backoff(attempt, max_retries, f"HTTP {code}")
                continue
            raise UploadError(f"HTTP {code} {reason}: {message}"[:300]) from e
        except RETRYABLE_EXCEPTIONS as e:
            attempt = backoff(attempt, max_retries, f"network error ({e.__class__.__name__})")
    if "id" not in response:
        raise UploadError(f"YouTube ne video id nahi di: {response}"[:300])
    log.info("   upload 100%")
    return str(response["id"])


# ---------------------------------------------------------------- commands

def cmd_dry_run(cfg: Config) -> None:
    """Kuch upload/save nahi; sirf dikhata hai kya hoga."""
    wb, ws = open_workbook(cfg, read_only=True)
    rows = read_rows(ws, header_map(ws))
    wb.close()
    overlay_pending(rows, cfg)
    ready, problems = split_rows(rows, cfg)
    left = quota_left(cfg, load_quota(cfg))
    log.info("===== DRY RUN (kuch upload nahi hoga, Excel nahi badlega) =====")
    log.info(f"Aaj ka quota bacha: {left}/{cfg.daily_limit}")
    log.info(f"Ready rows: {len(ready)}")
    for i, (r, when) in enumerate(ready, 1):
        mark = "  <- is run mein jayegi" if i <= left else ""
        log.info(f"  {i}. {when:%Y-%m-%d %H:%M} IST | {r.filename} | file mili | {r.title[:50]}{mark}")
    if problems:
        log.info(f"Problem wali rows (asli run mein FAILED hongi): {len(problems)}")
        for r, err in problems:
            log.info(f"  - {r.label}: {err}")
    if load_pending(cfg):
        log.info("Note: pending_updates.json mein purane update baaki hain - asli run unhe pehle Excel mein daalega.")


def cmd_status(cfg: Config) -> None:
    """Excel ka summary: kitni PENDING/SCHEDULED/FAILED aur agli 5 rows."""
    wb, ws = open_workbook(cfg, read_only=True)
    rows = read_rows(ws, header_map(ws))
    wb.close()
    overlay_pending(rows, cfg)
    counts: dict[str, int] = {}
    for r in rows:
        key = r.status or "PENDING"
        counts[key] = counts.get(key, 0) + 1
    log.info("===== STATUS =====")
    for key in ("PENDING", "SCHEDULED", "FAILED"):
        log.info(f"  {key:<10} {counts.pop(key, 0)}")
    for key, n in counts.items():
        log.info(f"  {key:<10} {n}")
    upcoming = sorted(
        ((r, w) for r in rows if is_candidate(r) and (w := parse_schedule(r.when_raw, cfg)) and w > now_ist(cfg)),
        key=lambda x: x[1],
    )[:5]
    log.info("Agli 5 YouTube rows:")
    for r, w in upcoming:
        log.info(f"  {w:%Y-%m-%d %H:%M} IST | {r.filename} | {r.title[:50]}")
    if not upcoming:
        log.info("  (koi nahi)")
    log.info(f"Aaj ka quota bacha: {quota_left(cfg, load_quota(cfg))}/{cfg.daily_limit}")


def open_for_write(cfg: Config) -> tuple[Workbook, Worksheet, dict[str, int]]:
    """Lock check + backup + pending updates lagakar Excel ko likhne ke liye kholta hai."""
    if not cfg.excel_path.exists():
        raise SystemExit(f"Excel file nahi mili: {cfg.excel_path}  (.env mein EXCEL_PATH check karo)")
    if excel_is_locked(cfg.excel_path):
        raise SystemExit(f"{cfg.excel_path.name} Excel me khuli hai, pehle band karo.")
    backup_excel(cfg)
    wb, ws = open_workbook(cfg, read_only=False)
    cols = header_map(ws)
    apply_pending(wb, ws, cols, cfg)
    return wb, ws, cols


def cmd_retry_failed(cfg: Config) -> None:
    """FAILED rows ko wapas PENDING karta hai (upload nahi karta)."""
    wb, ws, cols = open_for_write(cfg)
    failed = [r for r in read_rows(ws, cols) if r.status == "FAILED"]
    for r in failed:
        write_cells(ws, cols, r.excel_row, {"status": "PENDING", "log_message": "retry ke liye wapas PENDING"})
    if failed:
        wb.save(cfg.excel_path)
    log.info(f"{len(failed)} FAILED rows wapas PENDING kar di.")


def upload_one(youtube: Any, cfg: Config, row: Row, when: datetime) -> dict[str, str]:
    """Ek row upload karta hai aur Excel mein likhne wale updates lautata hai."""
    log.info(f"{row.label} | {when:%Y-%m-%d %H:%M} IST | {row.title[:60]}")
    try:
        video_id = upload_video(youtube, cfg.video_dir / row.filename, build_body(row, when, cfg), cfg.max_retries)
    except UploadError as e:
        log.error(f"   FAILED: {e}")
        return {"status": "FAILED", "log_message": str(e)}
    log.info(f"   SCHEDULED - video id {video_id}")
    return {
        "status": "SCHEDULED",
        "youtube_video_id": video_id,
        "processed_at": now_ist(cfg).strftime("%Y-%m-%d %H:%M:%S"),
        "log_message": f"YouTube scheduled for {when:%Y-%m-%d %H:%M} IST",
    }


def cmd_run(cfg: Config, max_n: int | None) -> None:
    """Asli run: problem rows FAILED, phir quota ke andar ready rows upload."""
    quota = load_quota(cfg)
    left = quota_left(cfg, quota)
    if left <= 0:
        log.info(f"Aaj ka quota khatam ({quota['uploads']}/{cfg.daily_limit}). Kal phir chalao.")
        return
    wb, ws, cols = open_for_write(cfg)
    ready, problems = split_rows(read_rows(ws, cols), cfg)
    for r, err in problems:
        log.warning(f"{r.label}: FAILED - {err}")
        save_row(wb, ws, cols, cfg, r, {"status": "FAILED", "log_message": err})
    allowed = min(left, max_n) if max_n else left
    batch = ready[:allowed]
    log.info(f"Ready: {len(ready)} | is run mein: {len(batch)} | aaj ka quota bacha: {left}/{cfg.daily_limit}")
    if not batch:
        log.info("Upload karne layak koi row nahi.")
        return
    youtube = get_youtube(cfg)
    done = 0
    for i, (row, when) in enumerate(batch):
        if i:
            pause = random.uniform(cfg.min_sleep, cfg.max_sleep)
            log.info(f"   {pause:.0f}s ruk raha hun...")
            time.sleep(pause)
        try:
            updates = upload_one(youtube, cfg, row, when)
        except QuotaExceeded as e:
            quota["uploads"] = cfg.daily_limit
            save_quota(cfg, quota)
            log.error(f"YouTube ne quota khatam bola ({e}). Ye row PENDING hi rahegi. Kal phir chalao.")
            break
        quota["uploads"] = int(quota["uploads"]) + 1
        save_quota(cfg, quota)
        save_row(wb, ws, cols, cfg, row, updates)
        done += updates["status"] == "SCHEDULED"
    log.info(f"Khatam: {done} scheduled. Aaj ka quota: {quota['uploads']}/{cfg.daily_limit}")
    if quota_left(cfg, quota) <= 0:
        log.info(f"Aaj ka quota khatam ({quota['uploads']}/{cfg.daily_limit}). Kal phir chalao.")
    if len(ready) > len(batch):
        log.info(f"{len(ready) - len(batch)} ready rows agle run ke liye bachi hain.")


def parse_args() -> argparse.Namespace:
    """Command line flags padhta hai."""
    p = argparse.ArgumentParser(description="theshivamiq YouTube Shorts scheduler")
    g = p.add_mutually_exclusive_group()
    g.add_argument("--dry-run", action="store_true", help="kuch upload nahi, sirf dikhao")
    g.add_argument("--status", action="store_true", help="Excel ka summary")
    g.add_argument("--retry-failed", action="store_true", help="FAILED rows ko PENDING karo")
    p.add_argument("--max", type=int, metavar="N", help=f"is run mein max N upload (1-{HARD_DAILY_CAP})")
    args = p.parse_args()
    if args.max is not None and not 1 <= args.max:
        p.error("--max kam se kam 1 hona chahiye")
    if args.max:
        args.max = min(args.max, HARD_DAILY_CAP)
    return args


def main() -> None:
    """Flags ke hisaab se sahi command chalata hai."""
    args = parse_args()
    cfg = load_config()
    setup_logging(cfg.log_file)
    log.info(f"----- scheduler start {datetime.now(cfg.tz):%Y-%m-%d %H:%M:%S} IST -----")
    try:
        if args.dry_run:
            cmd_dry_run(cfg)
        elif args.status:
            cmd_status(cfg)
        elif args.retry_failed:
            take_lock(cfg)
            try:
                cmd_retry_failed(cfg)
            finally:
                release_lock(cfg)
        else:
            take_lock(cfg)
            try:
                cmd_run(cfg, args.max)
            finally:
                release_lock(cfg)
    except SystemExit as e:
        if isinstance(e.code, str):
            log.error(e.code)
            raise SystemExit(1) from None
        raise
    except KeyboardInterrupt:
        log.warning("Aapne beech mein rok diya. Jo upload ho chuke the wo Excel mein save hain.")
    except Exception:
        log.exception("Anjaan galti - scheduler.log dekho / Claude ko bhejo.")
        raise SystemExit(1) from None


if __name__ == "__main__":
    main()
