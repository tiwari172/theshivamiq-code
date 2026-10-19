import sys, os, math, shutil
sys.path.insert(0, os.path.dirname(__file__))
from batch2_data import CASES, SHORTS  # nayi batch: batchN_data.py banao, yahan naam badlo
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

MASTER, OUT = sys.argv[1], sys.argv[2]
shutil.copy(MASTER, OUT)
wb = load_workbook(OUT)
M = load_workbook(MASTER)

STYLE = ("Flat 2D cartoon illustration, bold black outlines, same art style as the 'Detective Shivam' series, "
         "Indian setting, warm cinematic colours, no watermark. ")
GREEN = " Plain bright green (#00FF00) background, no text, no shadow."
SAFE = " Peaceful pose, YouTube-safe: no blood, no wounds."
FEMALE = {"Kavita", "Sunita", "Shobha", "Rashmi", "D'Souza", "Champa", "Devi", "Anjali", "Tanishka", "Munni", "Sneha", "Komal", "Lata",
          "Meenakshi", "Divya", "Mary", "Lakshmi", "Keerthi", "Jennifer", "Mariam", "Leela", "Sarla", "Payal", "Radha", "Sudha"}
HOOKS = ["SLAM_TITLE + SHAKE", "TYPE_HOOK", "SPLIT_FLAP + GLITCH"]
ENTRIES = ["SLIDE_R", "SILHOUETTE_REVEAL", "DOOR_LIGHT"]
ORD = {11: "ग्यारहवें", 12: "बारहवें", 13: "तेरहवें", 14: "चौदहवें", 15: "पंद्रहवें", 16: "सोलहवें", 17: "सत्रहवें", 18: "अठारहवें", 19: "उन्नीसवें", 20: "बीसवें"}
for c in CASES:
    c["id"] = f"A-{c['n']:03d}"; c["tag"] = f"A{c['n']:03d}"

def short(rn):
    t = [x.replace('.', '') for x in rn.split() if x.lower() not in ("ji", "sahab", "sir", "devi")]
    return t[-1]
def is_f(rn): return any(w in FEMALE for w in rn.replace('.', '').split())
def aid(tag, rn): return f"{tag}_{short(rn).upper().replace(chr(39), '')}"
def dur(vo, mn=1.8, wps=2.6, pad=0.35): return round(max(mn, len(vo.split()) / wps + pad), 1)
def ts(t): t = int(round(t)); return f"{t // 60}:{t % 60:02d}"

# ------------------------------------------------------------------ SCRIPT
ms = list(M["Script + Timing"].iter_rows(values_only=True))
intro = [list(r) for r in ms[1:] if r[0] == "INTRO"]
outro = [list(r) for r in ms[1:] if r[0] == "OUTRO"]
rows = []; t = 0.0; chapters = [(0, "Intro + niyam")]
def add(case, sec, vo, cap, screen="", asset="", pose="", motion="", sfx="", color="white", note="", d=None, extra=None):
    global t
    d = d if d is not None else (dur(vo) if vo else 2.0)
    rows.append(dict(r=[case, len(rows) + 1, sec, ts(t), ts(t + d), d, vo, cap, screen, asset, pose, motion, sfx, color, note], t0=t, d=d, x=extra or {}))
    t += d
for r in intro:
    add("INTRO", r[2], r[6], r[7], r[8], r[9], r[10], r[11], r[12], r[13], r[14] or "", d=float(r[5]), extra={"intro": True})

for i, c in enumerate(CASES):
    tag, N = c["tag"], len(c["suspects"])
    chapters.append((t, f"Case {c['id']}: {c['title']}"))
    add(tag, "Title card", "", f"CASE {c['id']}: {c['title'].upper()}", f"CASE {c['id']}\n{c['title'].upper()}\n[{c['level']}]", "CARD_TITLE", "",
        "title card slide-in + punch, level badge (Medium hara / Hard naranji / Very Hard laal)", "whoosh", "yellow", d=2.0, extra={"c": c, "i": i})
    add(tag, "Hook", c["hook"][0], c["hook"][1], c["hook"][2], c["bg"][0], "S2_17", "nopop + zoom, bada yellow sawaal", "heartbeat", "yellow", extra={"c": c, "i": i})
    add(tag, "Kahani", c["setup"][0][0], c["setup"][0][1], "", c["victim"][0], "", "night + heartbeat, dheere zoom (Ken Burns)" if c["night"] else "heartbeat, dheere zoom (Ken Burns)", "heartbeat", extra={"c": c, "i": i})
    add(tag, "Maut ka time", c["setup"][1][0], c["setup"][1][1], c["setup"][1][2], c["victim"][0], "", "laal text pop", "pop", "red", extra={"c": c, "i": i})
    add(tag, "Shivam entry", c["entry"][0], c["entry"][1], "", c["bg"][0], c["entry"][2], "AUTO_PUNCH, Shivam side se slide", "whoosh", extra={"c": c, "i": i})
    for k, s in enumerate(c["saboot"], 1):
        add(tag, f"Saboot {k}", "", "", "", s[0], s[2], "EVIDENCE_TRAY: sirf tasveer 1.5s, sab saboot SAME (clue alag NAHI)", "swish",
            note="ASLI CLUE (sirf aapke liye) | VO nahi (sirf tasveer)" if s[3] else "VO nahi (sirf tasveer)", d=1.5, extra={"c": c, "i": i, "k": k, "K": len(c["saboot"])})
    add(tag, "Bayaan shuru", "अब सुनिए सबके बयान।", "Ab suno sabke bayaan.", f"{N} SUSPECTS", c["bg"][0], "S2_05", "punch", "pop2", extra={"c": c, "i": i, "N": N})
    for k, s in enumerate(c["suspects"], 1):
        vo = s[0] + "। " + (s[1] + "। " if s[1] else "") + ("बोलीं... " if is_f(s[3]) else "बोले... ") + s[2]
        screen = f"SUSPECT {k}/{N}\n{s[3]}" + (f"\n({s[4]})" if s[4] else "")
        add(tag, f"Suspect {k}", vo, f"\"{s[5]}\"", screen, aid(tag, s[3]), "", "popup (rise / slide, baari-baari), naam yellow upar, quote neeche",
            "pop" if k % 2 else "pop2", "yellow", note="KILLER (sirf aapke liye)" if k == c["killer"] else "", extra={"c": c, "i": i, "k": k, "N": N})
    add(tag, "Galat ishara", c["galat"][0], c["galat"][1], "", aid(tag, c["suspects"][c["galat"][2] - 1][3]), "S2_03",
        "Shivam + us suspect side-by-side, halka zoom", "swish", extra={"c": c, "i": i})
    add(tag, "Soch (3 sec)", "", "Shivam SOCH\nraha hai...", "Shivam SOCH\nraha hai...", "", "S2_02", "zoom+music_off, +1s aur +2s par khaali music_off rows",
        "tick_fast", "yellow", d=3.0, extra={"c": c, "i": i})
    add(tag, "Pakad liya", "डिटेक्टिव शिवम ने कातिल पकड़ लिया। क्या आपने पकड़ा?", "Detective Shivam ne kaatil pakad liya. Kya aapne pakda?", "", "", "S2_06",
        "flash+music_off", "boom", extra={"c": c, "i": i})
    add(tag, "Sawaal", "कातिल कौन? वीडियो रोको, और कमेंट में लिखो केस नंबर, और कातिल का नाम।", f"KILLER KAUN? Pause karo! Comment: {c['id']} + naam",
        f"KILLER KAUN?\n{c['id']} + NAAM", "CARD_PAUSE", "S2_21 (na ho to S2_17)", "PAUSE icon + bada yellow text, 1 sec ruk kar", "ding", "yellow", extra={"c": c, "i": i, "N": N})
chapters.append((t, "Outro: jawab kab milega"))
for r in outro:
    add("OUTRO", r[2], r[6], r[7], r[8], r[9], r[10], r[11], r[12], r[13], r[14] or "", d=float(r[5]), extra={"outro": True})
TOTAL = t

# ------------------------------------------------------------------ MOTION PLAN (long)
mp = list(M["MOTION PLAN"].iter_rows(values_only=True))
m_intro = [list(r) for r in mp[1:] if r[0] == "INTRO"]
m_outro = [list(r) for r in mp[1:] if r[0] == "OUTRO"]
def chunks(text, sec):
    k = max(1, math.ceil(sec / 3.2))
    while k > 1 and sec / k < 1.8: k -= 1
    w = (text or "").split()
    if not w: return [("", sec / k)] * k
    per = math.ceil(len(w) / k)
    parts = [" ".join(w[j:j + per]) for j in range(0, len(w), per)]
    while len(parts) < k: parts.append("")
    return [(p, round(sec / k, 1)) for p in parts[:k]]
RULE_S = "Rule: har saboot SAME (clue wale ko bhi koi zoom/gola/peela NAHI)"
RULE_P = "Rule: har suspect SAME entry + SAME sfx pattern (killer ko alag mat dikhao)"
def motion_beats(R, L916=False):
    """Ek script row -> beats: (vo, sec, asset, pose, motion, camera, layout, screen, trans, sfx, music, note)"""
    r, x = R["r"], R["x"]; sec_name, d = r[2], R["d"]; c = x.get("c")
    cap = r[7] or ""; asset, pose, screen = r[9] or "", r[10] or "-", r[8] or ""
    B = []
    def b(vo, s, mo, cam, lay, scr="", tr="cut", sfx="", mus="normal", note=""):
        B.append([vo, s, asset or "-", pose or "-", mo, cam, lay, scr, tr, sfx, mus, note])
    ch = chunks(cap, d)
    if sec_name == "Title card":
        mo = "CASE_SWIPE -> FILE_OPEN" + (" + WARNING_STRIP" if c["level"] != "MEDIUM" else "")
        b(cap, d, mo, "static", "file beech mein, badge right-upar", screen, "whip-pan left", "whoosh",
          note=f"CASE_HUD yahin se shuru: 'CASE {c['id']} | {c['n'] - 10}/10'")
    elif sec_name == "Hook":
        hk = HOOKS[x["i"] % 3]
        for j, (vo, s) in enumerate(ch):
            if j == 0: b(vo, s, hk, "bg dheere andar 1.00->1.06", "sawaal upar 25-45%, Shivam BUST neeche (S2_17)" if L916 else "text beech, Shivam BUST right 1/3 (S2_17)", screen, "FLASH", "heartbeat")
            else: b(vo, s, "BUST punch-in on sawaal", "PUNCH 1.0->1.12 aakhri shabd par", "Shivam BUST neeche, text upar" if L916 else "Shivam BUST right")
    elif sec_name == "Kahani":
        for j, (vo, s) in enumerate(ch):
            if j == 0: b(vo, s, f"MAP_PIN chip '{c['city']}'", "bg drift left->right", "pin + chip upar-beech" if L916 else "pin bg ke upar beech", "", "mask wipe", "heartbeat")
            else: b(vo, s, "victim image + " + ("NIGHT + " if c["night"] else "") + "CRIME_TAPE", "Ken Burns dheere andar 1.00->1.08",
                    "victim beech 55% height" if L916 else "victim beech 70% height", "", "cross-dissolve 6 frame")
        if len(ch) == 1: B[-1][4] += " + victim image + " + ("NIGHT + " if c["night"] else "") + "CRIME_TAPE"
    elif sec_name == "Maut ka time":
        for j, (vo, s) in enumerate(ch):
            if j == 0: b(vo, s, "TIME_WINDOW + laal text pop", "PUNCH 1.0->1.12 time par", "ghadi upar-beech, time chip ghadi ke neeche" if L916 else "ghadi left 1/3, time chip right", screen, "cut", "pop")
            else: b(vo, s, "TIME_WINDOW arc bhar raha", "static", "ghadi upar-beech" if L916 else "ghadi left")
    elif sec_name == "Shivam entry":
        b(cap, d, ENTRIES[x["i"] % 3], "halka push 1.00->1.04", "Shivam beech-neeche 60% height, bg dhundhla" if L916 else "Shivam right 1/3, bg dhundhla", "", "whip-pan right", "whoosh")
    elif sec_name.startswith("Saboot"):
        b("", 1.5, f"EVIDENCE_TRAY tag 'SABOOT {x['k']}/{x['K']}'", "upar se desk view, static",
          "polaroid 90% chaudai beech" if L916 else "polaroid beech, Shivam S2 chhota left-neeche", "", "slide left", "swish", note=RULE_S)
    elif sec_name == "Bayaan shuru":
        b(cap, d, f"SUSPECT_GRID {x['N']} saaye + SLAM_TITLE", "static", "grid beech, number neeche", screen, "cut + AUTO_PUNCH", "pop2", note="Saaye mein koi alag rang/size nahi")
    elif sec_name.startswith("Suspect"):
        for j, (vo, s) in enumerate(ch):
            if j == 0: b(vo, s, "INTERROGATION + SUSPECT_CARD", "upar se halka tilt-down", "suspect upar-beech 50% height, card uske upar" if L916 else "suspect left 40%, card right-upar",
                         screen, "cut + AUTO_PUNCH", "pop" if x["k"] % 2 else "pop2", note=RULE_P)
            elif j == 1: b(vo, s, "STATEMENT_SLIP (bayaan type)", "push 1.00->1.03", "suspect upar, parchi neeche 55-70%" if L916 else "suspect left, parchi right", note=RULE_P)
            else: b(vo, s, "SLIP_FLIP (baaki bayaan)", "pan 3% right", "suspect upar, parchi neeche" if L916 else "suspect left, parchi right", note=RULE_P)
    elif sec_name == "Galat ishara":
        for j, (vo, s_) in enumerate(ch):
            if j == 0: b(vo, s_, "SPLIT_SCREEN Shivam | galat suspect", "dono taraf se andar", "Shivam upar, suspect neeche" if L916 else "Shivam left, suspect right", "", "mask wipe beech se", "swish")
            else: b(vo, s_, "DIVE + MARKER (sirf galat aadmi par)", "tez zoom galat suspect par", "suspect beech", "", "cut", "", note="DIVE/MARKER sirf galat-ishara wale par, killer par KABHI nahi")
    elif sec_name.startswith("Soch"):
        b("Shivam SOCH / raha hai...", 3.0, "SOCH_WATCH + PAUSE_HINT", "zoom 1.00->1.06", "Shivam beech S2_02, ⏸ upar-right (15% ke neeche)" if L916 else "Shivam beech S2_02, ⏸ right-upar",
          "Shivam SOCH / raha hai...", "cut", "tick_fast", "OFF", "3 sec chup VO; music_off 3 baar (har 1 sec)")
    elif sec_name == "Pakad liya":
        h = round(d / 2, 1)
        b("Detective Shivam ne kaatil", h, "BLACKOUT 0.3s -> TWIST_SLAM 'PAKAD LIYA!' + FLASH", "SHAKE tez + PUNCH 1.12", "Shivam S2_06 beech", "", "cut", "boom", "OFF")
        b("pakad liya. Kya aapne pakda?", h, "hold 'KYA AAPNE PAKDA?'", "static", "Shivam S2_06 beech")
    elif sec_name == "Sawaal":
        h = round(d / 2, 1)
        b("KILLER KAUN? Pause karo!", h, f"LINEUP_Q {x['N']} saaye", "dheere left->right pan", "line poori chaudai mein", screen, "cut", "ding", "DUCK", "Koi eliminate NAHI, killer alag nahi")
        b(f"Comment: {c['id']} + naam", h, f"COMMENT_TYPE '{c['id']} + NAAM' + ⏸", "static", "comment box beech, ⏸ left", "", "cut", "", "DUCK")
    return B

mrows = []
for r in m_intro: mrows.append(r)
for R in rows:
    r = R["r"]
    if r[0] in ("INTRO", "OUTRO"): continue
    tt = R["t0"]
    for j, bt in enumerate(motion_beats(R), 1):
        vo, s, asset, pose, mo, cam, lay, scr, tr, sfx, mus, note = bt
        mrows.append([R["x"]["c"]["id"], f"{r[1]}.{j}", round(tt, 1), round(tt + s, 1), s, r[2], vo, asset, pose, mo, cam, lay, scr, tr, sfx, mus, note])
        tt += s
off = TOTAL - sum(float(r[4]) for r in m_outro)
base = rows[-len(outro)]["r"][1]
k0 = None
for r in m_outro:
    rr = list(r); dd = float(rr[4]); rr[2] = round(off, 1); rr[3] = round(off + dd, 1); off += dd
    num, sub = str(rr[1]).split(".")
    if k0 is None: k0 = int(num)
    rr[1] = f"{base + int(num) - k0}.{sub}"
    mrows.append(rr)

# ------------------------------------------------------------------ SHORTS + REELS / SHORTS MOTION PLAN
srows, smrows = [], []
for sn, (cn, keep, hk) in enumerate(SHORTS, 1):
    c = [x for x in CASES if x["n"] == cn][0]; tag = c["tag"]; L = f"R{sn} {c['id']}"
    order = [f"Hook (short wala): {hk[0]} / {hk[1]}"]
    vo_lines = [f"[whispers] {hk[0]}", c["setup"][0][0].rstrip("।") + "!", "[serious] " + c["setup"][1][0], c["entry"][0], "[thoughtful] " + c["jagah"][0]]
    for k in keep:
        s = c["suspects"][k - 1]; vo_lines.append(f"{s[0]}... {s[7]} \"{s[2]}\"")
    vo_lines += ["⏸ 3 sec chup (Shivam soch)",
                 "[dramatic] डिटेक्टिव शिवम ने कातिल पकड़ लिया!  +  CTA_HI (आंसर कमेंट करो, सही आंसर कल पिन होगा।)  +  SIGNATURE_HI (शिवम की नज़र से कोई नहीं बचता।)"]
    order += ["Kahani", "Maut ka time", "Shivam entry"] + [f"Saboot {k}" for k in range(1, len(c["saboot"]) + 1)]
    order += [f"Suspect {k} -> SUSPECT {j}/{len(keep)}" for j, k in enumerate(keep, 1)] + ["Soch (3 sec)", "Pakad liya + CTA + Signature (purani files)"]
    # ---- shorts motion beats
    st = 0.0; bn = 0
    def sb(sec_name, beats):
        global st
        for j, bt in enumerate(beats, 1):
            vo, s, asset, pose, mo, cam, lay, scr, tr, sfx, mus, note = bt
            smrows.append([L, f"{bn}.{j}", round(st, 1), round(st + s, 1), s, sec_name, vo, asset, pose, mo, cam, lay, scr, tr, sfx, mus, note]); st += s
    ci = CASES.index(c)
    fake = lambda sec, cap, d, asset="", pose="", screen="", **kw: dict(r=[tag, 0, sec, "", "", d, "", cap, screen, asset, pose, "", "", "", ""], d=d, x=dict(c=c, i=ci, **kw))
    rd = lambda txt: dur(txt, 1.8, 2.9, 0.3)
    bn += 1; sb("Hook", motion_beats(fake("Hook", hk[1], rd(hk[0]), c["bg"][0], "S2_17", hk[2]), True))
    bn += 1; sb("Kahani", motion_beats(fake("Kahani", c["setup"][0][1], rd(c["setup"][0][0]), c["victim"][0]), True))
    bn += 1; sb("Maut ka time", motion_beats(fake("Maut ka time", c["setup"][1][1], rd(c["setup"][1][0]), c["victim"][0], "", c["setup"][1][2]), True))
    bn += 1; sb("Shivam entry", motion_beats(fake("Shivam entry", c["entry"][1], rd(c["entry"][0]), c["bg"][0], c["entry"][2]), True))
    bn += 1
    jd = max(rd(c["jagah"][0]), 1.1 * len(c["saboot"]) + 1.0); jparts = chunks(c["jagah"][1], jd)
    mon = [f"EVIDENCE_TRAY MONTAGE (saboot 1-{len(c['saboot'])}, har ek ~1.1s, SAB EK JAISE): polaroid girta -> 0.7s ruk -> slide-out", "MONTAGE jaari (same aana-jaana)",
           "MONTAGE jaari. AAKHRI: saare polaroid whip-pan se bahar"]
    sb("Saboot (montage)", [[vo, s, c["saboot"][0][0] + " ... " + c["saboot"][-1][0], "S2_04", mon[min(j, 2)], "upar se desk view, har polaroid par halka push",
                             "polaroid 90% chaudai beech", "", "slide up" if j == 0 else "cut", "swish" if j == 0 else "", "normal", RULE_S] for j, (vo, s) in enumerate(jparts)])
    for j, k in enumerate(keep, 1):
        s = c["suspects"][k - 1]; bn += 1
        R = fake(f"Suspect {j}/{len(keep)}", f"\"{s[5]}\"", rd(s[2]) + 0.6, aid(tag, s[3]), "-",
                 f"SUSPECT {j}/{len(keep)}\n{s[3]}" + (f"\n({s[4]})" if s[4] else ""), k=j, N=len(keep))
        R["r"][2] = "Suspect"; sb(f"Suspect {j}/{len(keep)}", motion_beats(R, True))
    bn += 1; sb("Soch (3 sec)", motion_beats(fake("Soch (3 sec)", "", 3.0, "", "S2_02"), True))
    bn += 1; sb("Pakad liya", motion_beats(fake("Pakad liya", "", 3.8, "", "S2_06"), True))
    bn += 1
    sb("End (CTA + signature)", [["Answer comment karo, sahi answer kal pin hoga.", 3.2, "-", "S2_17", f"COMMENT_TYPE '{c['id']} + NAAM' + 'SAHI ANSWER KAL PIN' chip", "PUNCH 1.12",
                                  "comment box beech 40-60%", "COMMENT KARO / KAL PIN HOGA", "whip-pan up", "ding", "DUCK", "Jawab kabhi nahi"],
                                 ["/ Shivam ki nazar se koi nahi bachta.", 3.1, "-", "S2_17", "END_SIGN", "dheere andar 1.00->1.05", "Shivam close-up beech", "", "cut", "", "DUCK", "Jawab kabhi nahi"]])
    srows.append([f"Short {sn}", f"Case {c['id']}: {c['title']} ({c['level']})", ", ".join(map(str, keep)), "\n".join(order),
                  "[dramatic] डिटेक्टिव शिवम ने कातिल पकड़ लिया!  +  CTA_HI (आंसर कमेंट करो, सही आंसर कल पिन होगा।)  +  SIGNATURE_HI (शिवम की नज़र से कोई नहीं बचता।)",
                  "Detective Shivam ne kaatil pakad liya!\nAnswer comment karo, sahi answer kal pin hoga.\nShivam ki nazar se koi nahi bachta.", f"~{int(st)} sec",
                  "Jawab short mein KABHI nahi. Killer + galat-ishara wala suspect hamesha rakho (card par). VO niyam: Start Here 'Reel VO niyam'.", "\n".join(vo_lines)])

# ------------------------------------------------------------------ IMAGE PROMPTS
irows = [list(r) for r in list(M["Image Prompts"].iter_rows(values_only=True))[1:] if str(r[0]).startswith(("BG_INTRO", "CARD_"))]
for c in CASES:
    tag = c["tag"]
    irows.append([c["bg"][0], tag, "Background", c["bg"][1], STYLE + c["bg"][1] + ". No people, no text. Make 16:9 (1920x1080) and 9:16 (1080x1920)."])
    irows.append([c["victim"][0], tag, "Victim", c["victim"][1], STYLE + c["victim"][1] + "." + SAFE + GREEN])
    for s in c["saboot"]:
        txt = " Text on the board/paper/screen must be clearly readable exactly as written." if "'" in s[1] else ""
        irows.append([s[0], tag, "Saboot (prop)", s[1], STYLE + s[1] + "." + txt + " Close-up object, centered." + GREEN])
    for s in c["suspects"]:
        irows.append([aid(tag, s[3]), tag, "Suspect", s[3] + (f" ({s[4]})" if s[4] else ""), STYLE + s[6] + ". Full body, standing, facing camera." + GREEN])

# ------------------------------------------------------------------ PINNED / UPLOAD / ANSWER
prow = [list(list(M["Pinned + Jawab"].iter_rows(values_only=True))[1])]
prow[0][2] = prow[0][2].replace("A-001 Naam, A-002 Naam ... A-010 Naam", "A-011 Naam, A-012 Naam ... A-020 Naam")
for c in CASES:
    killer = c["suspects"][c["killer"] - 1][3]
    rest = c["reveal"][1].split("! ", 1)[1]
    prow.append([c["id"], "Agle hafte ki video (jawab)", f"{ORD[c['n']]} केस का जवाब: " + c["reveal"][0], f"Case {c['id']}: " + c["reveal"][1],
                 f"🔍 Case {c['id']} ({c['title']}) ka jawab: {short(killer).upper()}!\n{rest}\n👍 Jisne sahi pakda tha, is comment ko LIKE karo!\nPoori explanation + winners: YouTube full video (link bio mein)"])
chap_txt = "\n".join(f"{ts(tt)} {name}" for tt, name in chapters)
mu = list(M["Upload Details"].iter_rows(values_only=True))
desc = [r[2] for r in mu if r[0] == "YouTube Long" and r[1] == "Description"][0]
a, b2 = desc.index("⏱ Chapters:\n"), desc.index("\n\nContest rules")
desc = (desc[:a] + "⏱ Chapters:\n" + chap_txt + desc[b2:]).replace("A-001 Naam, A-002 Naam ... A-010 Naam", "A-011 Naam, A-012 Naam ... A-020 Naam")
urows = [
 ["YouTube Long", "Title (option 1)", "10 Naye Murder Mysteries | 99% Log FAIL! Kya Aap Solve Kar Paoge? | Detective Shivam Batch 2"],
 ["YouTube Long", "Title (option 2)", "Saboot Saamne Hai, Kaatil Pakdo! 10 Cases (Part 2) | Detective Shivam"],
 ["YouTube Long", "Title (option 3)", "Sirf 1% Log Saare 10 Kaatil Pakad Paate | Hindi Mystery Quiz | Batch 2"],
 ["YouTube Long", "Description", desc],
] + [list(r) for r in mu[1:] if r[0] == "YouTube Long" and r[1] not in ("Title (option 1)", "Title (option 2)", "Title (option 3)", "Description")]
for sn, (cn, keep, hk) in enumerate(SHORTS, 1):
    c = [x for x in CASES if x["n"] == cn][0]
    urows += [
     [f"YouTube Short {sn}", "Title", f"Case {c['id']}: {c['title']} | {len(keep)} Suspects, 1 Kaatil 🕵️ #shorts"],
     [f"YouTube Short {sn}", "Description", f"Case {c['id']}: {c['title']}. {len(keep)} suspects, 1 kaatil. Kaatil kaun? Baaki suspects + POINTS sirf full video mein 👉 [long video link]\n#DetectiveShivam #MurderMystery #Paheli #shorts"],
     [f"YouTube Short {sn}", "Related video", "Short ke 'Related video' mein Batch 2 long video lagao (YouTube Studio > Short edit)"],
     [f"Instagram Reel {sn}", "Caption", f"🕵️ Case {c['id']}: {c['title']}\n{c['setup'][0][1]}\n{len(keep)} suspects, 1 kaatil. 99% log galat naam lete hain. Tum? Comment karo 👇\nBaaki suspects + POINTS wala muqabla YouTube full video mein (link bio mein)"],
     [f"Instagram Reel {sn}", "Hashtags", "#detectiveshivam #murdermystery #paheli #riddles #hindiriddles #mysteryreels #brainteaser #theshivamiq #crimepuzzle #reelsindia"],
     [f"Instagram Reel {sn}", "Cover text", "99% LOG FAIL! / KILLER KAUN?"],
     [f"Instagram Reel {sn}", "Pinned comment (post karte hi)", f"Case {c['id']} 🕵️ Kaatil ka naam comment karo! Jawab agle Sunday. Points wala muqabla YouTube full video par (link bio mein)."],
     [f"YouTube Short {sn}", "Pinned comment (post karte hi)", f"Case {c['id']} 🕵️ Kaatil kaun? Jawab agle Sunday. POINTS sirf full video ke comments par 👉 [long video link]"],
     [f"Shorts + Reel {sn}", "Jawab ke baad pinned", f"'Pinned + Jawab' sheet ka aakhri column (Case {c['id']}) — purana pinned hata kar ye lagao"],
     [f"Instagram Reel {sn}", "Story", "Reel share karo + YouTube link sticker: 'Full case + points yahan'"],
    ]
urows.append(["Release plan", "Order", "Sunday: long video (shuru mein Batch 1 ke jawab + winners) | Mon/Wed/Fri: Short 1/2/3 (YouTube + Insta) | Agla Sunday: Batch 2 ke jawab + Batch 3"])
arows = [[c["id"], c["title"], c["level"], c["killer"], c["suspects"][c["killer"] - 1][3], c["answer"] + f"  [Pattern: {c['pattern']}]", c["search"] or "-"] for c in CASES]

# ------------------------------------------------------------------ WRITE
thin = Side(style="thin", color="BBBBBB")
def refill(name, data, fills=None):
    ws = wb[name]
    hdr_h = ws.row_dimensions[1].height
    if ws.max_row > 1: ws.delete_rows(2, ws.max_row)
    for r in data: ws.append(r)
    for ri, row in enumerate(ws.iter_rows(min_row=2), 2):
        f = fills(row) if fills else None
        for cc in row:
            cc.alignment = Alignment(vertical="top", wrap_text=True)
            cc.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            cc.fill = PatternFill("solid", fgColor=f) if f else PatternFill(fill_type=None)
            cc.font = Font()
def sfill(row):
    note = str(row[14].value or "")
    if "KILLER" in note: return "FDE2E2"
    if "CLUE" in note: return "FFF4C2"
    if row[0].value in ("INTRO", "OUTRO"): return "E8F0FE"
refill("Script + Timing", [R["r"] for R in rows], sfill)
refill("MOTION PLAN", mrows)
refill("Shorts + Reels", srows)
refill("SHORTS MOTION PLAN", smrows)
refill("Image Prompts", irows, lambda r: "FFF4C2" if r[2].value == "Saboot (prop)" else None)
refill("Pinned + Jawab", prow)
refill("Upload Details", urows)
refill("Answer Key", arows, lambda r: "FDE2E2")
# Start Here
sh = wb["Start Here"]
for row in sh.iter_rows(min_row=2):
    k, v = row[0].value, row[1].value or ""
    v = v.replace("Batch 1 (YouTube Long Video 01)", "Batch 2 (YouTube Long Video 02)")
    if k == "Video": v = f"10 cases (5 Medium, 3 Hard, 2 Very Hard), A-011 se A-020. Andaazan lambai: {int(TOTAL // 60)} min {int(TOTAL % 60)} sec (asli timing voice banne ke baad)."
    if k == "Case ID": v = "Category A: A-011 se A-020 (Batch 2). Title, hook, caption mein sirf case ID + victim/jagah + suspense. Clue KABHI nahi."
    if k == "Comment format": v = "Har case ke end mein: 'A-011 Naam', 'A-012 Naam'... (Category A, case 011 se 020). Isse comment checker aasaani se points dega."
    row[1].value = v
# Cover example
for row in wb["Cover Design"].iter_rows(min_row=2):
    for cc in row:
        if isinstance(cc.value, str): cc.value = cc.value.replace("CASE A-001 / JUDGE SAHAB KI AAKHRI RAAT / [MEDIUM]", "CASE A-011 / BANK MANAGER KA BUNGALOW / [MEDIUM]")
wb.active = 0
wb.save(OUT)
print("script", len(rows), "motion", len(mrows), "shorts motion", len(smrows), "images", len(irows), "total", ts(TOTAL),
      "| shorts:", [r[6] for r in srows])
