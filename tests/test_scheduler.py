import json
import os
import time
from datetime import datetime, timedelta
from pathlib import Path
from unittest import mock
from zoneinfo import ZoneInfo

import pytest
from googleapiclient.errors import HttpError
from openpyxl import Workbook, load_workbook

import scheduler as sch

IST = ZoneInfo("Asia/Kolkata")


def make_cfg(tmp: Path, **kw) -> sch.Config:
    (tmp / "output").mkdir(exist_ok=True)
    base = dict(video_dir=tmp / "output", excel_path=tmp / "schedule.xlsx", sheet_name="Post Schedule", tz=IST,
                client_secrets=tmp / "client_secret.json", token_file=tmp / "token.json", category_id="24",
                daily_limit=6, min_sleep=0, max_sleep=0, max_retries=3, quota_file=tmp / "quota_state.json",
                pending_file=tmp / "pending_updates.json", log_file=tmp / "scheduler.log", backup_dir=tmp / "backups",
                lock_file=tmp / "scheduler.lock")
    base.update(kw)
    return sch.Config(**base)


def make_excel(cfg, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Post Schedule"
    ws.append(list(sch.NEEDED_COLUMNS) + ["ig_caption"])
    for r in rows:
        ws.append(r)
    wb.save(cfg.excel_path)


def future(h=5):
    return (datetime.now(IST) + timedelta(hours=h)).strftime("%Y-%m-%d %H:%M:%S")


def http_error(status, reason=""):
    resp = mock.Mock(status=status, reason="x")
    body = json.dumps({"error": {"errors": [{"reason": reason}], "message": "msg"}}).encode()
    return HttpError(resp, body)


# ---------------------------------------------------------------- quota date
def test_quota_day_is_pacific_time():
    pt = datetime.now(ZoneInfo("America/Los_Angeles")).strftime("%Y-%m-%d")
    assert sch.utc_today() == pt


def test_quota_state_format_same(tmp_path):
    cfg = make_cfg(tmp_path)
    st = sch.load_quota(cfg)
    st["uploads"] = 2
    sch.save_quota(cfg, st)
    assert set(json.loads(cfg.quota_file.read_text())) == {"date", "uploads"}
    assert sch.quota_left(cfg, sch.load_quota(cfg)) == 4


# ---------------------------------------------------------------- .env
def test_env_blank_values_use_default(monkeypatch):
    monkeypatch.setattr(sch, "load_dotenv", lambda *a, **k: None)
    monkeypatch.setenv("DAILY_UPLOAD_LIMIT", "")
    monkeypatch.setenv("MIN_SLEEP_SECONDS", " ")
    monkeypatch.delenv("YT_CONTAINS_SYNTHETIC_MEDIA", raising=False)
    cfg = sch.load_config()
    assert cfg.daily_limit == 6 and cfg.min_sleep == 5.0 and cfg.synthetic_media is None


def test_env_bad_number_clear_message(monkeypatch):
    monkeypatch.setattr(sch, "load_dotenv", lambda *a, **k: None)
    monkeypatch.setenv("MAX_RETRIES", "paanch")
    with pytest.raises(SystemExit, match="MAX_RETRIES"):
        sch.load_config()


# ---------------------------------------------------------------- lock
def test_lock_blocks_second_run(tmp_path):
    cfg = make_cfg(tmp_path)
    sch.take_lock(cfg)
    with pytest.raises(SystemExit, match="pehle se chal raha"):
        sch.take_lock(cfg)
    sch.release_lock(cfg)
    sch.take_lock(cfg)          # chhoot gaya to dobara mil jaata hai
    sch.release_lock(cfg)


def test_stale_lock_removed(tmp_path):
    cfg = make_cfg(tmp_path)
    cfg.lock_file.write_text("123")
    old = time.time() - 7 * 3600
    os.utime(cfg.lock_file, (old, old))
    sch.take_lock(cfg)
    assert cfg.lock_file.exists()
    sch.release_lock(cfg)


# ---------------------------------------------------------------- body
def test_body_synthetic_flag_only_when_set(tmp_path):
    row = sch.Row(2, "", "", "a.mp4", "Title <x>", "desc", "a, b", None, "YOUTUBE", "")
    when = datetime.now(IST) + timedelta(days=1)
    b = sch.build_body(row, when, make_cfg(tmp_path))
    assert "containsSyntheticMedia" not in b["status"] and b["snippet"]["title"] == "Title x"
    b = sch.build_body(row, when, make_cfg(tmp_path, synthetic_media=True))
    assert b["status"]["containsSyntheticMedia"] is True
    assert b["status"]["privacyStatus"] == "private" and b["status"]["publishAt"].endswith("Z")


# ---------------------------------------------------------------- upload retry (YouTube mock)
class FakeReq:
    def __init__(self, steps):
        self.steps = list(steps)

    def next_chunk(self):
        s = self.steps.pop(0)
        if isinstance(s, Exception):
            raise s
        return s


def fake_youtube(req):
    yt = mock.Mock()
    yt.videos.return_value.insert.return_value = req
    return yt


def test_upload_retries_then_succeeds(tmp_path, monkeypatch):
    monkeypatch.setattr(sch.time, "sleep", lambda s: None)
    monkeypatch.setattr(sch, "MediaFileUpload", mock.Mock())
    req = FakeReq([http_error(503), OSError("net"), (None, {"id": "VID1"})])
    assert sch.upload_video(fake_youtube(req), tmp_path / "a.mp4", {}, 3) == "VID1"


def test_upload_quota_and_permanent_errors(tmp_path, monkeypatch):
    monkeypatch.setattr(sch.time, "sleep", lambda s: None)
    monkeypatch.setattr(sch, "MediaFileUpload", mock.Mock())
    with pytest.raises(sch.QuotaExceeded):
        sch.upload_video(fake_youtube(FakeReq([http_error(403, "quotaExceeded")])), tmp_path / "a.mp4", {}, 3)
    with pytest.raises(sch.UploadError):
        sch.upload_video(fake_youtube(FakeReq([http_error(400, "invalidPublishAt")])), tmp_path / "a.mp4", {}, 3)
    with pytest.raises(sch.UploadError, match="baar try"):
        sch.upload_video(fake_youtube(FakeReq([http_error(500)] * 5)), tmp_path / "a.mp4", {}, 2)


# ---------------------------------------------------------------- poora run (Excel + YouTube mock)
def test_run_marks_scheduled_failed_and_keeps_other_columns(tmp_path, monkeypatch):
    cfg = make_cfg(tmp_path)
    (cfg.video_dir / "01 A.mp4").write_bytes(b"x")
    make_excel(cfg, [
        ["01 A.mp4", "Title A", "d", "t", future(), "YOUTUBE", "", "", "", "", "IG caption A"],
        ["02 B.mp4", "Title B", "d", "t", future(), "BOTH", "PENDING", "", "", "", "IG caption B"],
        [None] * 11,
    ])
    monkeypatch.setattr(sch, "get_youtube", lambda c: "YT")
    monkeypatch.setattr(sch, "upload_video", lambda yt, p, b, r: "VID_" + p.stem[:2])
    sch.cmd_run(cfg, None)
    ws = load_workbook(cfg.excel_path)["Post Schedule"]
    assert ws["G2"].value == "SCHEDULED" and ws["I2"].value == "VID_01"
    assert ws["G3"].value == "FAILED" and "nahi mili" in ws["H3"].value
    assert ws["K2"].value == "IG caption A"           # ig_caption ko chhua nahi
    assert json.loads(cfg.quota_file.read_text())["uploads"] == 1


def test_run_save_fail_goes_to_pending_then_applied(tmp_path, monkeypatch):
    cfg = make_cfg(tmp_path)
    (cfg.video_dir / "01 A.mp4").write_bytes(b"x")
    make_excel(cfg, [["01 A.mp4", "Title A", "d", "t", future(), "YOUTUBE", "", "", "", ""]])
    monkeypatch.setattr(sch, "get_youtube", lambda c: "YT")
    monkeypatch.setattr(sch, "upload_video", lambda *a: "VID1")
    real_save = sch.Workbook.save
    monkeypatch.setattr(sch.Workbook, "save", lambda self, p: (_ for _ in ()).throw(PermissionError("locked")))
    sch.cmd_run(cfg, None)
    assert json.loads(cfg.pending_file.read_text())[0]["updates"]["youtube_video_id"] == "VID1"
    monkeypatch.setattr(sch.Workbook, "save", real_save)
    sch.cmd_retry_failed(cfg)                           # koi bhi likhne wala run pending pehle daalta hai
    ws = load_workbook(cfg.excel_path)["Post Schedule"]
    assert ws["G2"].value == "SCHEDULED" and ws["I2"].value == "VID1"
    assert not cfg.pending_file.exists()


def test_quota_exceeded_keeps_row_pending(tmp_path, monkeypatch):
    cfg = make_cfg(tmp_path)
    (cfg.video_dir / "01 A.mp4").write_bytes(b"x")
    make_excel(cfg, [["01 A.mp4", "Title A", "d", "t", future(), "YOUTUBE", "", "", "", ""]])
    monkeypatch.setattr(sch, "get_youtube", lambda c: "YT")

    def boom(*a):
        raise sch.QuotaExceeded("quotaExceeded")
    monkeypatch.setattr(sch, "upload_video", boom)
    sch.cmd_run(cfg, None)
    ws = load_workbook(cfg.excel_path)["Post Schedule"]
    assert ws["G2"].value in (None, "")
    assert json.loads(cfg.quota_file.read_text())["uploads"] == 6
