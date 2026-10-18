"""theshivamiq Comment Checker
Roz chalao: py checker.py            (saari reels)
Test:       py checker.py --case 17-002   (sirf ek case, leaderboard_TEST.xlsx banti hai)
"""
import os, sys, re, json, argparse
from datetime import datetime, timedelta, timezone
from collections import Counter, defaultdict
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

import ig_api, textmatch as tm, answer_key as akey

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "data")
COMMENTS_FILE = os.path.join(DATA, "comments.json")
STATE_FILE = os.path.join(DATA, "state.json")
MY_USER = "theshivamiq"
IST = timezone(timedelta(hours=5, minutes=30))
NOW = datetime.now(timezone.utc)

SAHI, GALAT, SHAK, NONE = "SAHI", "GALAT", "SHAK", "-"
CUTOFF_HOURS = 24      # reel post ke itne ghante tak ke comment gine jaate hain

# ---------------- POINTS (yahin badlo) ----------------
PTS_SAHI = 10          # har sahi jawab
PTS_KOSHISH = 1        # galat / shak – koshish ke liye
WATCH_SEC = 60         # post ke 1 min baad speed ki ghadi chalu (8:00 ki reel -> 8:01 se); pehle = bina dekhe andaaza
# bonus kabhi sahi jawab (10) se zyada nahi – warna log jaldi ke chakkar mein andaaza lagayenge
SPEED = [(30, 10), (60, 8), (5 * 60, 6), (15 * 60, 4), (30 * 60, 2), (60 * 60, 1)]   # (reel ke baad kitne sec tak, points)
TOP3_PTS = [10, 7, 5]  # reel par pehle 3 sahi jawab
STREAK_EVERY, STREAK_PTS = 5, 10    # har 5 lagatar reels par comment
DAY_STREAK = {5: 15, 7: 25, 15: 50, 30: 100, 90: 250, 180: 500, 365: 1000}   # lagatar din comment kiya
PTS_VERIFIED = 50      # Telegram + YouTube + Instagram follow (ek baar)
PTS_STORY = 15         # story mein @theshivamiq mention (din mein 1)
POINTS_START = datetime(2026, 10, 19, tzinfo=timezone(timedelta(hours=5, minutes=30)))   # points yahin se (Somvaar); pehle ke 0
POINTS_DOW_FROM = "2026-10-19"   # is hafte se Detective of the Week POINTS se; pehle ke hafte SAHI ginti se (reel mein yahi bataya)


def speed_points(sec):
    after = sec - WATCH_SEC
    if after < 0:
        return 0
    for lim, p in SPEED:
        if after <= lim:
            return p
    return 0


def human_secs(sec):
    sec = int(sec)
    if sec < 60:
        return "%d sec" % sec
    if sec < 3600:
        return "%d min" % (sec // 60)
    return "%d ghante" % (sec // 3600)


def log(*a):
    print(*a, flush=True)


def jload(p, default):
    try:
        with open(p, encoding="utf-8") as f:
            return json.load(f)
    except ValueError:
        # file kharab hai: chupchap {} se overwrite karne se pehle uski copy bacha lo (data na khoye)
        bad = p + ".kharab-" + datetime.now().strftime("%Y%m%d_%H%M%S")
        try:
            os.replace(p, bad)
            log("[!] %s kharab thi, copy rakh di: %s" % (os.path.basename(p), os.path.basename(bad)))
        except OSError:
            pass
        return default
    except OSError:
        return default


def jsave(p, obj):
    os.makedirs(DATA, exist_ok=True)
    tmp = p + ".new"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False)
    os.replace(tmp, p)


def ts(s):
    """Instagram time ('2026-10-17T10:15:00+0000') -> datetime"""
    return datetime.strptime(s[:19], "%Y-%m-%dT%H:%M:%S").replace(tzinfo=timezone.utc)


def ist(dt):
    return dt.astimezone(IST).strftime("%Y-%m-%d %H:%M") if dt else ""


def parse_user_time(v):
    if not v:
        return None
    if isinstance(v, datetime):
        return v.replace(tzinfo=IST) if v.tzinfo is None else v
    for f in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y %H:%M", "%d/%m/%Y %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(v).strip(), f).replace(tzinfo=IST)
        except ValueError:
            pass
    return None


# ---------------------------------------------------------------- Instagram se data
# "Case NF-001 Part 2" – Part 2 case number ke THEEK baad hona chahiye
# (taaki Part 1 ke caption mein "Part 2 kal aayega" likha ho to wo Part 2 na maani jaaye)
PART2_RE = re.compile(r"case\s*#?\s*[A-Za-z0-9]+\s*-\s*\d+\s*[|,:\-–—(\[]?\s*"
                      r"(?:part|pt\.?|भाग|पार्ट)\s*[-#:]?\s*(?:2|ii|two|२)\b", re.I)


def case_of(caption, when=None):
    """Caption se case: "Case 22-001" (hamesha) ya "Case NF-001" jaise naye ID (sirf POINTS_START se)."""
    m = re.search(r"[Cc]ase\s*#?\s*(\d+)\s*-\s*(\d+)", caption or "")
    if m:
        return "%02d-%03d" % (int(m.group(1)), int(m.group(2)))
    if when is not None and when < POINTS_START:
        return None                       # purani reels ki ginti waisi hi rahe
    m = re.search(r"[Cc]ase\s*#?\s*([A-Za-z][A-Za-z0-9]{0,5})\s*-\s*(\d{1,4})\b", caption or "")
    return "%s-%03d" % (m.group(1).upper(), int(m.group(2))) if m else None


def is_part2(caption, when):
    """Part 2 = jawab wali reel (19 Oct se). Iske comment nahi ginte, aur iske post hote hi Part 1 ki ginti band."""
    return when >= POINTS_START and bool(PART2_RE.search(caption or ""))


def case_fixes():
    """reel_case_fix.txt: caption mein galat case ho to sahi case."""
    out = {}
    try:
        for line in open(os.path.join(HERE, "reel_case_fix.txt"), encoding="utf-8-sig"):
            line = line.split("#")[0].strip()
            if "=" in line:
                a, b = [x.strip() for x in line.split("=", 1)]
                code = re.sub(r".*/(reel|p)/", "", a).strip("/ ")
                c = case_of("case " + b)
                if code and c:
                    out[code] = c
    except OSError:
        pass
    return out


def fetch_media():
    items = ig_api.get_all("me/media", {
        "fields": "id,caption,timestamp,permalink,media_product_type,comments_count,like_count", "limit": 50})
    out, fix = [], case_fixes()
    for m in items:
        code = re.sub(r".*/(reel|p)/", "", m.get("permalink") or "").strip("/")
        when = ts(m["timestamp"]) if m.get("timestamp") else NOW
        c = fix.get(code) or case_of(m.get("caption"), when)
        if c:
            m["case"] = c
            m["part2"] = is_part2(m.get("caption"), when)
            out.append(m)
    return out


def fetch_comments(media_id):
    rows = {}
    top = ig_api.get_all(media_id + "/comments", {
        "fields": "id,text,timestamp,from{id,username},replies{id,text,timestamp,from{id,username}}", "limit": 50})

    def row(c, parent):
        fr = c.get("from") or {}
        return {"u": fr.get("username") or c.get("username", ""), "uid": fr.get("id", ""),
                "t": c.get("text", ""), "ts": c["timestamp"], "p": parent}
    for c in top:
        rows[c["id"]] = row(c, "")
        for r in (c.get("replies") or {}).get("data", []):
            rows[r["id"]] = row(r, c["id"])
    return rows


def fetch_insights(media_id):
    out = {}
    for met in ("views,reach,likes,comments,shares,saved", "ig_reels_avg_watch_time"):
        try:
            r = ig_api.get(media_id + "/insights", {"metric": met})
            out.update({d["name"]: d["values"][0]["value"] for d in r.get("data", [])})
        except ig_api.IGError:
            pass
    return out


# ---------------------------------------------------------------- answer samajhna
class Case:
    def __init__(self, case, row):
        self.case = case
        self.answer = str(row.get("Sahi Answer") or "").strip()
        self.quiz = None
        q = tm.quiz_answers(self.answer)
        if len(q) >= 2 and re.fullmatch(r"[\s\d\.\)\-:a-dA-D,]+", self.answer):
            self.quiz = q
        self.suspects = akey.parse_suspects(row.get("Suspects (number naam)"))
        try:
            self.number = int(row.get("Sahi Number") or 0)
        except (TypeError, ValueError):
            self.number = 0
        extra = [x.strip() for x in str(row.get("Doosri Spelling") or "").split(",") if x.strip()]
        if not self.number:  # suspects mein se khud dhoondo
            ak = {tm.key(w) for w in tm.tokens(self.answer) if w not in tm.STOP}
            for n, lab in self.suspects:
                if ak & {tm.key(w) for w in tm.tokens(lab)}:
                    self.number = n
                    break
        self.max_n = max([n for n, _ in self.suspects] + [self.number, 0])
        # screen par 1-2-3 card nahi tha (answers_extra "number nahi") -> comment ke number answer nahi
        self.use_numbers = not row.get("No Numbers")
        # har suspect ke pehchaan wale shabd
        alias = defaultdict(set)
        for n, lab in self.suspects:
            for part in lab.split("/"):
                for w in tm.tokens(part):
                    if w not in tm.STOP and len(w) >= 3 and not w.isdigit():
                        alias[n].add(tm.key(w))
        if self.number:
            for part in [self.answer] + extra:
                for w in tm.tokens(part):
                    if w not in tm.STOP and len(w) >= 3 and not w.isdigit():
                        alias[self.number].add(tm.key(w))
        # jo shabd do suspects mein common hai (jaise "aunty") wo kisi ko pehchaanta nahi
        cnt = Counter(k for s in alias.values() for k in s)
        self.alias = {n: {k for k in s if cnt[k] == 1} for n, s in alias.items()}
        self.pin = parse_user_time(row.get("Answer Pin Time"))
        self.pin_auto = False

    def judge(self, text):
        """-> (result, score_text)"""
        text = re.sub(r"@[\w.]+", " ", text or "")   # @username ke number/naam answer nahi
        if self.quiz:
            got = tm.quiz_answers(text)
            if not got:
                return NONE, ""
            ok = sum(1 for q, a in self.quiz.items() if got.get(q) == a)
            return (SAHI if ok == len(self.quiz) else GALAT), "%d/%d" % (ok, len(self.quiz))
        # "Nikhil kyuki Pintu ne..." -> kyuki se pehle wala hissa hi answer hai
        parts = REASON_RE.split(text, 1)
        if len(parts) > 1:
            first = self._hits(parts[0])
            if first:
                return self._verdict(first)
        return self._verdict(self._hits(text))

    def _hits(self, text):
        strong, weak = tm.numbers_detail(text, self.max_n) if self.use_numbers else ([], [])
        hits = set()
        for w in tm.tokens(text):
            if w in tm.STOP or w.isdigit():
                continue
            k = tm.key(w)
            found = [n for n, keys in self.alias.items() if any(tm.similar(k, a) for a in keys)]
            if len(found) == 1:
                hits.add(found[0])
        # naam mila to beech wale (reason ke) number chhod do; naam nahi mila to wahi number answer
        hits |= set(strong) if hits else set(strong or weak)
        return hits

    def _verdict(self, hits):
        if not hits:
            return NONE, ""
        if hits == {self.number}:
            return SAHI, ""
        if self.number in hits:
            return SHAK, ""
        return GALAT, ""


REASON_RE = re.compile(r"\b(?:kyuki|kyunki|kyoki|kyonki|kyuke|kiuki|kiunki|qki|qk|kyu ki|kyun ki|because|bcoz|bcz|"
                       r"bcs|bc|coz|cause|since|as|reason|isliye|kyunke)\b|क्योंकि|क्यूंकि|क्युकी|:-|\s[-–—]+\s", re.I)


def norm_text(t):
    return " ".join(tm.tokens(t))


# ---------------------------------------------------------------- Excel helpers
HEAD = PatternFill("solid", fgColor="1F3864")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
GREEN = PatternFill("solid", fgColor="E2EFDA")
RED = PatternFill("solid", fgColor="FCE4D6")


def sheet(wb, title, cols, widths):
    ws = wb.create_sheet(title)
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        c = ws.cell(1, i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "A2"
    return ws


def load_verified():
    """verified.txt: jinhone YouTube subscribe + Telegram join kiya (ek line mein ek Instagram username)."""
    out = set()
    try:
        for line in open(os.path.join(HERE, "verified.txt"), encoding="utf-8-sig"):
            line = line.split("#")[0].strip().lstrip("@").lower()
            if line:
                out.add(line)
    except OSError:
        pass
    return out


def load_verified_dates():
    """{username: verify ka din} – bot line ke aage '# bot YYYY-MM-DD' likhta hai; na ho to aaj."""
    out = {}
    try:
        for line in open(os.path.join(HERE, "verified.txt"), encoding="utf-8-sig"):
            nm = line.split("#")[0].strip().lstrip("@").lower()
            if not nm:
                continue
            m = re.search(r"(\d{4}-\d{2}-\d{2})", line)
            out[nm] = datetime.strptime(m.group(1), "%Y-%m-%d").replace(hour=12, tzinfo=IST) if m else NOW
    except OSError:
        pass
    return out


def read_shak_decisions(path):
    """Shak sheet mein user ne jo Sahi/Galat likha."""
    out = {}
    if not os.path.exists(path):
        return out
    try:
        wb = load_workbook(path, read_only=True)
    except Exception:
        return out
    if "Shak wale Comments" not in wb.sheetnames:
        return out
    ws = wb["Shak wale Comments"]
    rows = ws.iter_rows(values_only=True)
    hdr = [str(h or "") for h in next(rows, [])]
    try:
        fi, ii = hdr.index("Faisla (Sahi/Galat)"), hdr.index("Comment ID")
    except ValueError:
        return out
    for r in rows:
        if not r or len(r) <= max(fi, ii):
            continue
        f = str(r[fi] or "").strip().lower()
        cid = str(r[ii] or "").strip()
        if cid and f in ("sahi", "galat", "s", "g", "right", "wrong", "yes", "no"):
            out[cid] = SAHI if f in ("sahi", "s", "right", "yes") else GALAT
    wb.close()
    return out


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--case", help="sirf ye case (test)")
    ap.add_argument("--offline", action="store_true", help="Instagram se naya data mat lao")
    args = ap.parse_args()
    out_file = os.path.join(HERE, "leaderboard_TEST.xlsx" if args.case else "leaderboard.xlsx")

    store = jload(COMMENTS_FILE, {})
    state = jload(STATE_FILE, {})
    state.setdefault("decisions", {})
    state.setdefault("memory", {})
    state.setdefault("rank_history", {})

    # 0) user ke Shak wale faisle yaad karo (Excel abhi likhne se pehle)
    newdec = read_shak_decisions(out_file)
    state["decisions"].update(newdec)

    if not args.offline:
        # 1) token refresh (hafte mein ek baar)
        last = state.get("token_refreshed")
        if not last or NOW - datetime.fromisoformat(last) > timedelta(days=7):
            try:
                days = ig_api.refresh_token()
                state["token_refreshed"] = NOW.isoformat()
                log("Token refresh ho gaya (%s din aur chalega)." % days)
            except ig_api.IGError as e:
                log("Token refresh abhi nahi hua (koi baat nahi):", e)
        # 2) reels
        log("Reels la raha hoon...")
        media = fetch_media()
        for m in media:
            s = store.setdefault(m["id"], {"comments": {}})
            s.update({k: m.get(k) for k in ("case", "timestamp", "permalink", "comments_count", "like_count", "part2")})
        media = [m for m in media if not m.get("part2")]   # jawab wali reel ke comment laane ki zaroorat nahi
        if args.case:
            media = [m for m in media if m["case"] == args.case]
            if not media:
                log("Case %s ki koi reel nahi mili." % args.case)
                return 1
        # 3) comments (purani reels jinka data pin ke baad aa chuka, unhe dobara nahi)
        failed, ok_n, fail_run = [], 0, 0
        for i, m in enumerate(media, 1):
            s = store[m["id"]]
            age = NOW - ts(m["timestamp"])
            last_f = s.get("fetched_at")
            if last_f and age > timedelta(days=10) and not args.case and \
                    datetime.fromisoformat(last_f) - ts(m["timestamp"]) > timedelta(days=10):
                continue
            log("  [%d/%d] %s  comments la raha hoon (%s)..." % (i, len(media), m["case"], m.get("comments_count")))
            try:
                s["comments"].update(fetch_comments(m["id"]))
                ok_n, fail_run = ok_n + 1, 0
            except ig_api.IGError as e:
                # ek reel ki network/limit galti se poora leaderboard na ruke: purane comments se kaam chalao
                failed.append(m["case"])
                fail_run += 1
                log("  [!] %s ke comments nahi aaye (%s) – purana data use hoga, kal dobara try" % (m["case"], e))
                if fail_run >= 2:
                    if not ok_n:
                        raise   # ek bhi reel nahi aayi = token/internet ki problem, admin ko PROBLEM dikhe
                    log("  [!] Lagatar 2 reels fail (Instagram limit?) – baaki reels purane data se")
                    break
                continue
            s["fetched_at"] = NOW.isoformat()
            if age < timedelta(days=30) or "insights" not in s:
                s["insights"] = fetch_insights(m["id"])
            jsave(COMMENTS_FILE, store)
        if failed:
            log("  [!] In reels ke comments is baar nahi aaye: %s" % ", ".join(failed))

    # 4) answer key (nayi cases khud judti hain)
    links = {}
    for mid, s in store.items():
        if s.get("case") and not s.get("part2"):
            links.setdefault(s["case"], s.get("permalink"))
    try:
        added = akey.add_missing(links)
        if added:
            log("Answer key mein nayi cases judi:", ", ".join(added))
    except PermissionError:
        log("answer_key.xlsx khuli hai, nayi cases nahi jud payi (Excel band karke dobara chalao).")
    keyrows = akey.load()

    # Part 2 (jawab wali reel) kab aayi: uske baad Part 1 par comment nahi ginte
    reveal = {}
    for mid, s in store.items():
        if s.get("case") and s.get("part2"):
            t = ts(s["timestamp"])
            reveal[s["case"]] = min(t, reveal.get(s["case"], t))
    reels = sorted([(mid, s) for mid, s in store.items() if s.get("case") and not s.get("part2")],
                   key=lambda x: x[1]["timestamp"])
    if args.case:
        reels = [r for r in reels if r[1]["case"] == args.case]

    cases, cutoff = {}, {}
    for mid, s in reels:
        c = s["case"]
        if c not in cases:
            cases[c] = Case(c, keyrows.get(c, {}))
        # reel post hone ke 24 ghante tak ke comment hi ginte hain (answer 24 ghante mein pin hota hai).
        # answer_key mein "Answer Pin Time" haath se bhara ho to wahi.
        cutoff[mid] = cases[c].pin or (ts(s["timestamp"]) + timedelta(hours=CUTOFF_HOURS))
        if c in reveal:                    # Part 2 (jawab) pehle aa gaya to wahin tak
            cutoff[mid] = min(cutoff[mid], reveal[c])

    # 5) har comment check
    per_case_user = {}          # (case,user) -> pehla jawab
    commented = defaultdict(set)  # user -> media ids
    udays = defaultdict(set)      # user -> kin dino (IST) comment kiya
    names = {}
    case_stats = defaultdict(Counter)
    num_votes = defaultdict(Counter)
    for mid, s in reels:
        cs = cases[s["case"]]
        for cid, x in sorted(s["comments"].items(), key=lambda kv: kv[1]["ts"]):
            if (x["u"] or "").lower() == MY_USER:
                continue
            u = x.get("uid") or (x["u"] or "").lower()
            if not u:
                continue
            commented[u].add(mid)
            udays[u].add(ts(x["ts"]).astimezone(IST).date())
            names[u] = x["u"]
            t = ts(x["ts"])
            if t > cutoff[mid]:
                continue
            if (cs.case, u) in per_case_user:
                continue
            res, score = cs.judge(x["t"])
            nt = norm_text(x["t"])
            mem = state["memory"].get(cs.case, {})
            how = "program"
            if cid in state["decisions"]:
                res, how = state["decisions"][cid], "aapka faisla"
                if res in (SAHI, GALAT):
                    state["memory"].setdefault(cs.case, {})[nt] = res
            elif res == SHAK and nt in mem:
                res, how = mem[nt], "yaad se"
            if res == NONE:
                continue  # "nice", emoji – jawab nahi, agla comment dekho
            nums = tm.numbers(x["t"], cs.max_n)
            if len(nums) == 1:
                num_votes[cs.case][nums[0]] += 1
            per_case_user[(cs.case, u)] = {"case": cs.case, "u": x["u"], "uk": u, "text": x["t"], "time": t, "res": res,
                                           "score": score, "cid": cid, "mid": mid, "how": how,
                                           "link": s.get("permalink")}
        case_stats[cs.case]["comments"] += sum(1 for x in s["comments"].values() if x["u"] and x["u"].lower() != MY_USER)
    for a in per_case_user.values():
        case_stats[a["case"]][a["res"]] += 1

    # 6) leaderboard
    # hafta = Somvaar 00:00 se Ravivaar raat tak (IST). "week" = pichhla poora hafta (Detective of the Week isi se,
    # poore hafte same rehta hai), "live" = is hafte abhi tak. "month" = pichhla poora mahina.
    now_i = NOW.astimezone(IST)
    this_mon = (now_i - timedelta(days=now_i.weekday())).replace(hour=0, minute=0, second=0, microsecond=0)
    prev_mon = this_mon - timedelta(days=7)
    this_m1 = now_i.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    prev_m1 = (this_m1 - timedelta(days=1)).replace(day=1)
    week_label = "%s – %s" % (prev_mon.strftime("%d %b"), (this_mon - timedelta(days=1)).strftime("%d %b"))
    month_label = prev_m1.strftime("%B %Y")
    users = defaultdict(lambda: {"name": "", "sahi": 0, "week": 0, "live": 0, "month": 0, "first": None, "answers": 0})
    for a in per_case_user.values():
        d = users[a["uk"]]
        d["name"] = a["u"]
        d["answers"] += 1
        if a["res"] == SAHI:
            d["sahi"] += 1
            if prev_mon <= a["time"] < this_mon:
                d["week"] += 1
            if a["time"] >= this_mon:
                d["live"] += 1
            if prev_m1 <= a["time"] < this_m1:
                d["month"] += 1
            if not d["first"] or a["time"] < d["first"]:
                d["first"] = a["time"]

    # streak: lagatar kitni case reels par comment kiya (24 ghante se nayi reel par comment na ho to streak nahi tootti)
    order = [mid for mid, s in reels]
    fresh = {mid for mid, s in reels if NOW - ts(s["timestamp"]) < timedelta(hours=24)}
    streak = {}
    for u, mids in commented.items():
        cur, best, run, broken = 0, 0, 0, False
        for mid in reversed(order):
            if mid in mids:
                run += 1
                if not broken:
                    cur = run
            elif mid in fresh and not broken:
                continue
            else:
                broken = True
                best = max(best, run)
                run = 0
        best = max(best, run)
        streak[u] = (cur, best, len(mids))

    JUNIOR, SENIOR = 3, 10   # kitne sahi par rank
    verified = load_verified()

    # 6b) POINTS – har point kab mila (time) taaki hafte / mahine ke points alag gine ja sakein
    post_t = {mid: ts(s["timestamp"]) for mid, s in reels}
    pts_log = []                       # (user, time, points, kya, case)
    sahi_by_mid = defaultdict(list)
    for a in per_case_user.values():
        a["sec"] = (a["time"] - post_t[a["mid"]]).total_seconds()
        if a["res"] == SAHI:
            sahi_by_mid[a["mid"]].append(a)
    place = {}                         # reel ke pehle 3 sahi (reel dekhne ke baad wale)
    for mid, lst in sahi_by_mid.items():
        ok = sorted((a for a in lst if a["sec"] >= WATCH_SEC), key=lambda a: a["time"])
        for i, a in enumerate(ok[:3]):
            place[a["cid"]] = i
    for a in per_case_user.values():
        a["pts"] = 0
        if a["res"] == SAHI:
            parts = [(PTS_SAHI, "Sahi jawab")]
            sp = speed_points(a["sec"])
            if sp:
                parts.append((sp, "Speed bonus (%s)" % human_secs(a["sec"])))
            if a["cid"] in place:
                i = place[a["cid"]]
                parts.append((TOP3_PTS[i], "Reel par %s sahi jawab" % ("PEHLA", "DOOSRA", "TEESRA")[i]))
        else:
            parts = [(PTS_KOSHISH, "Koshish (jawab diya)")]
        for p, why in parts:
            pts_log.append((a["uk"], a["time"], p, why, a["case"]))
            a["pts"] += p
    # streak bonus: har 5 lagatar reels par
    for u, mids in commented.items():
        run = 0
        for mid in order:
            if post_t[mid] < POINTS_START:
                continue
            if mid in mids:
                run += 1
                if run % STREAK_EVERY == 0:
                    pts_log.append((u, post_t[mid] + timedelta(hours=1), STREAK_PTS, "Streak %d reels 🔥" % run, ""))
            elif mid not in fresh:
                run = 0
    # verified (Telegram + YouTube + Instagram follow): ek baar
    name2u = {names.get(u, "").lower(): u for u in commented}
    vdates = load_verified_dates()
    for nm, day in vdates.items():
        u = name2u.get(nm)
        if u:
            pts_log.append((u, max(day, POINTS_START), PTS_VERIFIED, "Telegram + YouTube + Instagram follow", ""))
    # story mention (data/mentions.json: [[username, iso time], ...])
    seen_day = set()
    for nm, t in jload(os.path.join(DATA, "mentions.json"), []):
        u = name2u.get(str(nm).lower().lstrip("@"))
        try:
            tt = datetime.fromisoformat(t)
        except (TypeError, ValueError):
            continue
        key = (u, tt.astimezone(IST).date())
        if u and key not in seen_day:
            seen_day.add(key)
            pts_log.append((u, tt, PTS_STORY, "Story mein mention", ""))
    # din wali streak: jin dino reel aayi, un sab din comment kiya (jis din reel nahi aayi wo din nahi tootta)
    post_days = sorted({t.astimezone(IST).date() for t in post_t.values() if t >= POINTS_START})
    today_d = NOW.astimezone(IST).date()
    dstreak = {}
    for u, days in udays.items():
        run, best = 0, 0
        for d in post_days:
            if d in days:
                run += 1
                best = max(best, run)
                if run in DAY_STREAK:
                    pts_log.append((u, datetime(d.year, d.month, d.day, 23, 0, tzinfo=IST), DAY_STREAK[run],
                                    "%d din lagatar 🔥" % run, ""))
            elif d != today_d:
                run = 0
        dstreak[u] = (run, best)
    pts_log = [x for x in pts_log if x[1] >= POINTS_START]
    for u, t, p, why, case in pts_log:
        d = users.setdefault(u, {"name": names.get(u, u), "sahi": 0, "week": 0, "live": 0, "month": 0, "first": None, "answers": 0})
        d["pts"] = d.get("pts", 0) + p
        if prev_mon <= t < this_mon:
            d["pts_week"] = d.get("pts_week", 0) + p
        if t >= this_mon:
            d["pts_live"] = d.get("pts_live", 0) + p
        if prev_m1 <= t < this_m1:
            d["pts_month"] = d.get("pts_month", 0) + p
    for d in users.values():
        for k in ("pts", "pts_week", "pts_live", "pts_month"):
            d.setdefault(k, 0)

    def is_ok(u):
        return users[u]["name"].lower() in verified

    def pick(field, exclude=(), only_ok=True):
        c = [(d[field], -(d["first"].timestamp() if d["first"] else 0), u) for u, d in users.items()
             if d[field] > 0 and u not in exclude and (is_ok(u) or not only_ok)]
        return max(c)[2] if c else None

    def rank_table(only_ok):
        partner = pick("sahi", only_ok=only_ok)
        if partner and users[partner]["sahi"] < SENIOR:
            partner = None
        week = pick("pts_week" if prev_mon.strftime("%Y-%m-%d") >= POINTS_DOW_FROM else "week",
                    exclude=(partner,), only_ok=only_ok)
        out = {}
        for u, d in users.items():
            if only_ok and not is_ok(u):
                continue
            if u == partner:
                out[u] = "Shivam ka Partner"
            elif u == week:
                out[u] = "Detective of the Week"
            elif d["sahi"] >= SENIOR:
                out[u] = "Senior Detective"
            elif d["sahi"] >= JUNIOR:
                out[u] = "Junior Detective"
        return out, partner, week

    # rank sirf unhe jo verified.txt mein hain (YouTube subscribe + Telegram join)
    ranks, partner, week = rank_table(True)
    would, _, _ = rank_table(False)
    if not args.case:   # Telegram bot ke /rank ke liye
        order_sahi = sorted(users, key=lambda u: (-users[u]["pts"], -users[u]["sahi"]))
        jsave(os.path.join(DATA, "ranks.json"), {
            "updated": ist(NOW), "total_cases": len({s["case"] for _, s in reels}),
            "users": {users[u]["name"].lower(): {
                "name": users[u]["name"], "sahi": users[u]["sahi"], "week": users[u]["week"], "live": users[u]["live"],
                "jawab": users[u]["answers"], "pos": i + 1, "rank": ranks.get(u, ""), "would": would.get(u, ""),
                "streak": streak.get(u, (0, 0, 0))[0], "best_streak": streak.get(u, (0, 0, 0))[1],
                "pts": users[u]["pts"], "pts_week": users[u]["pts_week"], "pts_live": users[u]["pts_live"],
                "day_streak": dstreak.get(u, (0, 0))[0], "best_day_streak": dstreak.get(u, (0, 0))[1]}
                for i, u in enumerate(order_sahi)}})
        # Somvaar / mahine ki report (Telegram bot admin ko bhejta hai)
        def top(field, sfield, n=10):
            us = sorted((u for u in users if users[u][field] > 0),
                        key=lambda u: (-users[u][field], -users[u][sfield], users[u]["first"] or NOW))[:n]
            return [[users[u]["name"], users[u][field], users[u][sfield], is_ok(u), dstreak.get(u, (0, 0))[0]] for u in us]
        _, p_any, w_any = rank_table(False)
        jsave(os.path.join(DATA, "report.json"), {
            "made": ist(NOW), "week_label": week_label, "month_label": month_label,
            "week_top": top("pts_week", "week"), "month_top": top("pts_month", "month"), "all_top": top("pts", "sahi"),
            "dow_verified": users[week]["name"] if week else "", "dow_any": users[w_any]["name"] if w_any else "",
            "partner_verified": users[partner]["name"] if partner else "", "partner_any": users[p_any]["name"] if p_any else "",
            "seniors": sum(1 for d in users.values() if d["sahi"] >= SENIOR),
            "juniors": sum(1 for d in users.values() if JUNIOR <= d["sahi"] < SENIOR),
            "week_jawab": sum(1 for a in per_case_user.values() if prev_mon <= a["time"] < this_mon),
            "verified": len(verified)})
    RANK_ORDER = {"Junior Detective": 1, "Senior Detective": 2, "Detective of the Week": 3, "Shivam ka Partner": 4}

    # aaj ka shoutout: pichhli baar (kal ya pehle) se rank badha
    today = NOW.astimezone(IST).strftime("%Y-%m-%d")
    hist = state["rank_history"] if not args.case else {}
    prev_days = sorted(d for d in hist if d < today)
    prev = hist[prev_days[-1]] if prev_days else None
    shout = []
    for u, r in ranks.items():
        old = (prev or {}).get(u)
        if prev is None:
            if RANK_ORDER[r] >= 2:
                shout.append((u, "", r))
        elif RANK_ORDER[r] > RANK_ORDER.get(old, 0):
            shout.append((u, old or "", r))
    if not args.case:
        hist[today] = ranks
        for d in sorted(hist)[:-30]:
            del hist[d]

    # 7) Excel likho
    wb = Workbook()
    wb.remove(wb.active)
    ws = sheet(wb, "Leaderboard", ["#", "Username", "KUL POINTS", "Pichhle hafte points (" + week_label + ")",
                                   "Is hafte points abhi tak", "Kul sahi", "Pichhle hafte sahi", "Rank", "Din streak 🔥",
                                   "Sabse lambi din streak", "Reel streak", "Kul jawab", "Verified (YT+TG)", "Pehla sahi kab",
                                   "Pichhle mahine points (" + month_label + ")"],
               [5, 26, 11, 16, 13, 9, 11, 34, 10, 12, 10, 9, 11, 17, 14])
    lb = sorted(users.items(), key=lambda kv: (-kv[1]["pts"], -kv[1]["sahi"], kv[1]["first"] or NOW))
    n = 0
    for u, d in lb:
        if not d["pts"] and not d["sahi"]:
            continue
        n += 1
        r = ranks.get(u) or (("(verify baaki) " + would[u]) if u in would else "")
        ds = dstreak.get(u, (0, 0))
        ws.append([n, d["name"], d["pts"], d["pts_week"], d["pts_live"], d["sahi"], d["week"], r, ds[0], ds[1],
                   streak.get(u, (0,))[0], d["answers"], "Haan" if is_ok(u) else "", ist(d["first"]), d["pts_month"]])
        if ranks.get(u) in ("Shivam ka Partner", "Detective of the Week"):
            for c in ws[ws.max_row]:
                c.fill = YELLOW

    ws = sheet(wb, "Points ka hisaab", ["Username", "Case No.", "Kya", "Points", "Kab"], [26, 9, 40, 8, 17])
    for u, t, p, why, case in sorted(pts_log, key=lambda x: x[1], reverse=True):
        ws.append([names.get(u) or users[u]["name"], case, why, p, ist(t)])

    ws = sheet(wb, "Points ke niyam", ["Kaam", "Points"], [60, 10])
    ws.append(["Sahi jawab", PTS_SAHI])
    ws.append(["Galat / shak jawab (koshish)", PTS_KOSHISH])
    ws.append(["Speed: reel post ke %d sec baad ghadi chalu (pehle = bina dekhe, bonus nahi)" % WATCH_SEC, ""])
    for lim, p in SPEED:
        ws.append(["  Speed bonus: reel ke baad %s ke andar sahi" % human_secs(lim), p])
    for i, p in enumerate(TOP3_PTS):
        ws.append(["Reel par %s sahi jawab" % ("PEHLA", "DOOSRA", "TEESRA")[i], p])
    ws.append(["Har %d lagatar reels par comment" % STREAK_EVERY, STREAK_PTS])
    for dn, p in DAY_STREAK.items():
        ws.append(["%d din lagatar comment" % dn, p])
    ws.append(["Telegram join + YouTube subscribe + Instagram follow (ek baar)", PTS_VERIFIED])
    ws.append(["Story mein @theshivamiq mention (din mein 1)", PTS_STORY])

    ws = sheet(wb, "Aaj ka Shoutout", ["Username", "Pehle rank", "Naya rank", "Kul sahi"], [26, 22, 22, 10])
    shout.sort(key=lambda x: (-RANK_ORDER[x[2]], -users[x[0]]["sahi"]))
    for u, old, new in shout:
        ws.append([users[u]["name"], old or ("(pehli baar)" if prev is None else "-"), new, users[u]["sahi"]])
    if not shout:
        ws.append(["Aaj koi naya rank nahi"])

    ws = sheet(wb, "Pakka Fan (Streak)", ["Username", "Abhi ki streak", "Sabse lambi streak", "Kitni reels par comment",
                                          "Kul case reels", "Kul sahi"], [26, 14, 16, 16, 14, 10])
    fans = sorted(streak.items(), key=lambda kv: (-kv[1][0], -kv[1][1], -kv[1][2]))
    for u, (cur, best, cnt) in fans:
        if cur < 3:
            continue
        name = names.get(u, u)
        ws.append([name, cur, best, cnt, len(order), users[u]["sahi"] if u in users else 0])
        if cur == len(order):
            for c in ws[ws.max_row]:
                c.fill = GREEN

    ws = sheet(wb, "Shak wale Comments", ["Case No.", "Username", "Comment", "Time", "Faisla (Sahi/Galat)",
                                          "Kisne tay kiya", "Reel Link", "Comment ID"], [9, 22, 50, 17, 16, 14, 40, 20])
    dv = DataValidation(type="list", formula1='"Sahi,Galat"', allow_blank=True)
    ws.add_data_validation(dv)
    shak_rows = [a for a in per_case_user.values() if a["res"] == SHAK or a["how"] != "program"]
    shak_rows.sort(key=lambda a: (a["how"] != "program" and a["res"] != SHAK, a["case"], a["time"]))
    for a in shak_rows:
        fz = "" if a["res"] == SHAK else ("Sahi" if a["res"] == SAHI else "Galat")
        ws.append([a["case"], a["u"], a["text"], ist(a["time"]), fz, a["how"] if fz else "", a["link"], a["cid"]])
        ws.cell(ws.max_row, 3).alignment = Alignment(wrap_text=True, vertical="top")
        dv.add(ws.cell(ws.max_row, 5))
        if not fz:
            ws.cell(ws.max_row, 5).fill = YELLOW

    ws = sheet(wb, "Case Summary", ["Case No.", "Post date", "Kul comments", "Jawab dene wale", "Sahi", "Galat", "Shak",
                                    "Sahi %", "Sahi Number", "Sabse zyada likha number", "Kab tak ke comment gine", "Reel Link", "Dhyaan do"],
               [9, 12, 11, 12, 8, 8, 8, 8, 10, 14, 17, 40, 40])
    seen = set()
    for mid, s in reversed(reels):
        c = s["case"]
        if c in seen:
            continue
        seen.add(c)
        cs, st = cases[c], case_stats[c]
        tot = st[SAHI] + st[GALAT] + st[SHAK]
        top = num_votes[c].most_common(1)
        topn = top[0][0] if top else ""
        warn = []
        if not cs.answer and not cs.quiz:
            warn.append("answer_key mein Sahi Answer bharo")
        if cs.number and topn and topn != cs.number and top[0][1] >= 10:
            warn.append("log zyada %s likh rahe hain – Sahi Number check karo" % topn)
        if sum(1 for _, x in reels if x["case"] == c) > 1:
            warn.append("is case ki 2 reels hain")
        ws.append([c, ist(ts(s["timestamp"]))[:10], st["comments"], tot, st[SAHI], st[GALAT], st[SHAK],
                   ("%d%%" % round(100 * st[SAHI] / tot)) if tot else "", cs.number or "", topn,
                   ist(cutoff[mid]) + ("" if cs.pin else " (24 ghante)"), s.get("permalink"), "; ".join(warn)])
        if warn:
            ws.cell(ws.max_row, 13).fill = RED

    ws = sheet(wb, "Reel Report", ["Case No.", "Post date", "Views", "Reach", "Likes", "Comments", "Shares", "Saves",
                                   "Avg watch (sec)", "Sahi %", "Reel Link"], [9, 12, 10, 10, 8, 10, 8, 8, 12, 8, 40])
    for mid, s in reversed(reels):
        ins = s.get("insights") or {}
        st = case_stats[s["case"]]
        tot = st[SAHI] + st[GALAT] + st[SHAK]
        aw = ins.get("ig_reels_avg_watch_time")
        ws.append([s["case"], ist(ts(s["timestamp"]))[:10], ins.get("views"), ins.get("reach"),
                   ins.get("likes", s.get("like_count")), ins.get("comments", s.get("comments_count")),
                   ins.get("shares"), ins.get("saved"), round(aw / 1000, 1) if aw else "",
                   ("%d%%" % round(100 * st[SAHI] / tot)) if tot else "", s.get("permalink")])

    ws = sheet(wb, "Sabhi Jawab", ["Case No.", "Username", "Comment", "Time", "Result", "Kisne tay kiya"],
               [9, 22, 60, 17, 8, 14])
    for a in sorted(per_case_user.values(), key=lambda a: (a["case"], a["time"])):
        ws.append([a["case"], a["u"], a["text"], ist(a["time"]), a["res"], a["how"]])
        ws.cell(ws.max_row, 5).fill = GREEN if a["res"] == SAHI else RED if a["res"] == GALAT else YELLOW

    while True:
        try:
            wb.save(out_file)
            break
        except PermissionError:
            log("\n!! %s Excel mein khuli hai. Use band karo, phir Enter dabao..." % os.path.basename(out_file))
            try:
                input()
            except EOFError:
                return 1
    jsave(STATE_FILE, state)
    if not args.case:
        jsave(os.path.join(DATA, "shak_pending.json"), [
            {"cid": a["cid"], "case": a["case"], "u": a["u"], "text": a["text"][:300], "link": a["link"],
             "answer": cases[a["case"]].answer, "number": cases[a["case"]].number}
            for a in sorted(per_case_user.values(), key=lambda a: a["time"], reverse=True) if a["res"] == SHAK])

    # 8) chhota summary
    log("\n========== RESULT ==========")
    for c in sorted(cases, reverse=True)[:12]:
        st = case_stats[c]
        tot = st[SAHI] + st[GALAT] + st[SHAK]
        log("  %s : jawab %4d | sahi %4d | galat %4d | shak %3d | sahi %s" % (
            c, tot, st[SAHI], st[GALAT], st[SHAK], ("%d%%" % round(100 * st[SAHI] / tot)) if tot else "-"))
    if partner:
        log("  Shivam ka Partner: %s (%d sahi)" % (users[partner]["name"], users[partner]["sahi"]))
    if week:
        log("  Detective of the Week: %s (%d sahi, hafta %s)" % (users[week]["name"], users[week]["week"], week_label))
    log("  Aaj ke shoutout: %d log | Shak wale (faisla baaki): %d" % (
        len(shout), sum(1 for a in per_case_user.values() if a["res"] == SHAK)))
    log("\nFile: " + out_file)
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except ig_api.IGError as e:
        log("\nPROBLEM:", e)
        sys.exit(1)
