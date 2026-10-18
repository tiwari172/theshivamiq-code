"""answer_key.xlsx: har case ka sahi answer + suspects. Script (SCRIPT_*.txt) se khud banta hai.
User jo row badal de, program use kabhi overwrite nahi karta."""
import os, re, glob
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
import textmatch as tm

HERE = os.path.dirname(os.path.abspath(__file__))
KEY_FILE = os.path.join(HERE, "answer_key.xlsx")
STORY_ROOT = r"C:\Users\Hp\Desktop\reels story\category"
if not os.path.isdir(STORY_ROOT):          # server par: sync kiye hue scripts
    STORY_ROOT = os.path.join(HERE, "scripts")

COLS = ["Case No.", "Reel Link", "Sahi Answer", "Sahi Number", "Suspects (number naam)",
        "Doosri Spelling", "Answer Pin Time", "Note"]
WIDTH = [10, 45, 18, 12, 55, 30, 20, 40]


def find_script(case):
    cat, story = case.split("-", 1)
    if not (cat.isdigit() and story.isdigit()):   # naye ID (jaise NF-001) – script haath se / answer haath se
        return None
    folder = os.path.join(STORY_ROOT, "Category %d" % int(cat), "Story %02d" % int(story))
    files = glob.glob(os.path.join(folder, "SCRIPT_*.txt"))
    return files[0] if files else None


def _section(txt, name_re):
    m = re.search(r"=+\s*" + name_re + r"[^\n]*=+\n(.*?)(?=\n=+ |\Z)", txt, re.S | re.I)
    return m.group(1) if m else ""


_SAY = r"(?:bola|boli|bolin|bole|bolta|kehta|kaha|ne kaha|chillaya|बोला|बोली|बोले|बोलीं|चिल्लाया|चिल्लाई)"


_NOT_LAST = set("ke ka ki ko ne se me mein tha thi the hai hain poocha bataya kaha aaya aayi gaya liya diya".split())


def _label(raw):
    raw = re.sub(r"\b" + _SAY + r"\b", " ", raw, flags=re.I)
    raw = re.sub(r"\s+(ne|ji ne)\s*$", "", raw.strip(), flags=re.I)
    raw = re.sub(r"^\s*(aur|phir|tab|fir|and)\s+", "", raw.strip(), flags=re.I)
    words = re.findall(r"[A-Za-z]+|[ऀ-ॿ]+", raw)
    # naam/role hamesha bade akshar se (Driver, Sharma Aunty) – "rule tha", "sirf itna" jaise hisse nahi
    if not words or len(words) > 6:
        return ""
    latin = [w for w in words if re.match(r"[A-Za-z]", w)]
    if latin and (not any(w[0].isupper() for w in latin) or words[-1].lower() in _NOT_LAST):
        return ""
    label = " ".join(words[-3:])
    return re.sub(r"^(ka|ki|ke)\s+", "", label, flags=re.I)


def _suspects(body):
    body = re.sub(r"\[[^\]]*\]", " ", body)
    out = []
    # har bayan (quote) se pehle ka hissa = suspect ka label
    for m in re.finditer(r'([^."!?\n:]{1,80}?)\s*' + _SAY + r'?\s*(?:\.\.\.|…|:|,|—)?\s*"', body):
        if body[:m.end(1)].count('"') % 2 == 1:   # quote ke andar ka text
            continue
        raw = m.group(1)
        label = _label(raw) or _label(raw.split(",")[0])   # "Pati Nikhil, jo office se aaya, bola"
        if not label or "shivam" in label.lower():
            continue
        if any(tm.key(tm.translit(label)) == tm.key(tm.translit(x)) for x in out):
            continue
        out.append(label)
    return out


def parse_script(path):
    txt = open(path, encoding="utf-8-sig").read()
    suspects = []
    for sec in (r"CAPTION", r"(?:MAIN )?VOICEOVER"):
        suspects = _suspects(_section(txt, sec))
        if len(suspects) >= 2:
            break
    # sahi answer: PINNED "Answer: X –" warna ANSWER section ka pehla hissa
    ans = ""
    pin = _section(txt, r"PINNED COMMENT")
    m = re.search(r"Answer\s*:\s*([^–—\-\n.(]+)", pin)
    if m:
        ans = m.group(1).strip()
    if not ans:
        a = _section(txt, r"ANSWER(?! REVEAL)")
        a = re.sub(r"^\s*(Chor|Killer|Kaatil|Qatil|Culprit|Jawab|Answer)\s*(tha|thi|hai)?\s*:\s*", "", a.strip(), flags=re.I)
        ans = re.split(r"[.–—\n(]", a, 1)[0].strip()
    num = ""
    ak = {tm.key(w) for w in tm.tokens(ans) if w not in tm.STOP}
    best = 0
    for i, s in enumerate(suspects, 1):
        sk = {tm.key(w) for w in tm.tokens(s) if w not in tm.STOP}
        hit = len(ak & sk)
        if hit > best:
            best, num = hit, i
    return {"answer": ans, "number": num, "suspects": suspects}


def fmt_suspects(sus):
    return " | ".join("%d %s" % (i, s) for i, s in enumerate(sus, 1))


def parse_suspects(cell):
    out = []
    for part in str(cell or "").split("|"):
        part = part.strip()
        m = re.match(r"(\d+)\s*[.):-]?\s*(.+)", part)
        if m:
            out.append((int(m.group(1)), m.group(2).strip()))
        elif part:
            out.append((len(out) + 1, part))
    return out


def load():
    """{case: dict} – user ki likhi file + answers_extra.txt (jahan Excel mein answer khaali ho)."""
    rows = {}
    if os.path.exists(KEY_FILE):
        wb = load_workbook(KEY_FILE)
        ws = wb.active
        hdr = [str(c.value or "").strip() for c in ws[1]]
        for r in ws.iter_rows(min_row=2, values_only=True):
            d = dict(zip(hdr, r))
            case = str(d.get("Case No.") or "").strip()
            if not case:
                continue
            rows[case] = d
    for case, d in load_extra().items():
        if not str((rows.get(case) or {}).get("Sahi Answer") or "").strip():
            rows[case] = d                  # Excel mein haath se likha answer hamesha upar
    return rows


# ------------------------------------------------------------ answers_extra.txt (bina Excel)
# Ek line = ek case:   SS-003 | Neha | Riya, Rohit, Kunal, Bhola, Neha, Ramu | number nahi
#   2nd = sahi answer, 3rd = suspects (comma se), 4th = "number nahi" (screen par 1-2-3 card nahi tha,
#   to comment ka "3" kisi suspect ka number nahi maana jaata). '#' = note.
EXTRA_FILE = os.path.join(HERE, "answers_extra.txt")
SS_ROOT = os.path.join(os.path.dirname(STORY_ROOT), "Sab Pe Shak")


def load_extra():
    out = {}
    try:
        lines = open(EXTRA_FILE, encoding="utf-8-sig").read().splitlines()
    except OSError:
        return out
    for line in lines:
        line = line.split("#")[0].strip()
        parts = [p.strip() for p in line.split("|")]
        if len(parts) < 2 or not parts[0] or not parts[1]:
            continue
        sus = [s.strip() for s in (parts[2] if len(parts) > 2 else "").split(",") if s.strip()]
        out[parts[0].upper()] = {
            "Case No.": parts[0].upper(), "Sahi Answer": parts[1], "Sahi Number": "",
            "Suspects (number naam)": fmt_suspects(sus), "Doosri Spelling": "", "Answer Pin Time": "",
            "No Numbers": len(parts) > 3 and "nahi" in parts[3].lower()}
    return out


def ss_answer(series_folder):
    """Sab Pe Shak case folder -> (answer, suspects). Part 2 caption: '...pakad liya... bhaiya Rohit!'"""
    p2 = glob.glob(os.path.join(series_folder, "Part 2", "CAPTION_*_P2.txt"))
    p1 = glob.glob(os.path.join(series_folder, "Part 1", "CAPTION_*_P1.txt"))
    if not p2:
        return None, []
    txt = open(p2[0], encoding="utf-8-sig").read()
    m = re.search(r"pakad liya[.…\s]*([^!\n]+)!", txt, re.I)
    if not m:
        return None, []
    ans = m.group(1).split()[-1].strip(" .,")
    sus = []
    if p1:
        sus = [s for s in _suspects(open(p1[0], encoding="utf-8-sig").read()) if not s.lower().startswith("dr")]
    if not any(tm.key(tm.translit(ans)) in {tm.key(w) for w in tm.tokens(s)} for s in sus):
        sus.append(ans)
    return ans, sus


def ss_scan(root=None):
    """Sab Pe Shak ke saare case -> answers_extra.txt mein jo nahi hain wo jodo. Return: nayi cases."""
    root = root or SS_ROOT
    have = load_extra()
    new = []
    for folder in sorted(glob.glob(os.path.join(root, "SS-*"))):
        case = os.path.basename(folder).split()[0].upper()
        if case in have:
            continue
        ans, sus = ss_answer(folder)
        if ans:
            new.append("%s | %s | %s | number nahi" % (case, ans, ", ".join(sus)))
    if new:
        old = open(EXTRA_FILE, "rb").read() if os.path.exists(EXTRA_FILE) else b""
        with open(EXTRA_FILE, "a", encoding="utf-8") as f:
            if not old:
                f.write("# case | sahi answer | suspects (comma) | number nahi = screen par 1-2-3 card nahi tha\n")
            elif not old.endswith(b"\n"):
                f.write("\n")
            f.write("\n".join(new) + "\n")
    return new


def _style(ws):
    for i, (c, w) in enumerate(zip(COLS, WIDTH), 1):
        cell = ws.cell(1, i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
        ws.column_dimensions[cell.column_letter].width = w
    ws.freeze_panes = "A2"


def add_missing(case_links):
    """case_links: {case: reel_link}. Nayi cases ki row jodo; purani rows (user edits) chhuo mat,
    sirf khaali Reel Link bhar do. Return: nayi cases ki list."""
    if os.path.exists(KEY_FILE):
        wb = load_workbook(KEY_FILE)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Answer Key"
        ws.append(COLS)
    _style(ws)
    hdr = [str(c.value or "").strip() for c in ws[1]]
    ci = {h: i + 1 for i, h in enumerate(hdr)}
    have = {}
    for r in range(2, ws.max_row + 1):
        v = str(ws.cell(r, ci["Case No."]).value or "").strip()
        if v:
            have[v] = r
    added = []
    for case in sorted(case_links):
        link = case_links[case]
        if case in have:
            c = ws.cell(have[case], ci["Reel Link"])
            if not c.value:
                c.value = link
            continue
        path = find_script(case)
        if path:
            p = parse_script(path)
            note = ""
            if not p["number"]:
                note = "Sahi Number khud nahi mila – check karo"
            elif len(p["suspects"]) < 2:
                note = "Suspects kam mile – check karo"
            row = [case, link, p["answer"], p["number"], fmt_suspects(p["suspects"]), "", "", note]
        else:
            row = [case, link, "", "", "", "", "", "SCRIPT nahi mila – answer haath se bharo"]
        ws.append(row)
        for i in range(1, len(COLS) + 1):
            ws.cell(ws.max_row, i).alignment = Alignment(vertical="top", wrap_text=(i in (5, 8)))
        added.append(case)
    wb.save(KEY_FILE)
    return added


if __name__ == "__main__":
    import sys
    if sys.argv[1:2] == ["ss"]:          # py answer_key.py ss  -> Sab Pe Shak ke answers answers_extra.txt mein
        added = ss_scan()
        print("Nayi cases:" if added else "Koi nayi SS case nahi (sab pehle se answers_extra.txt mein).")
        for line in added:
            print("  " + line)
        sys.exit(0)
    for f in sorted(glob.glob(os.path.join(STORY_ROOT, "Category *", "Story *", "SCRIPT_*.txt")))[: int(sys.argv[1]) if len(sys.argv) > 1 else 9999]:
        p = parse_script(f)
        print(os.path.basename(f), "|", p["answer"], "|", p["number"], "|", fmt_suspects(p["suspects"]))
